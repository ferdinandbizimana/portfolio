"""RBC accountability: hand-checked figures on fake data, report wording, round trip."""

import io
from datetime import date
from pathlib import Path
from typing import Any

import pandas as pd
import pyarrow.parquet as pq
import pytest
from openpyxl import load_workbook

from app import datasets
from app.analyses.common.programmes import normalise_category
from app.analyses.rbc_accountability import (
    ANALYSIS,
    INCREASE,
    NO_CHANGE,
    REDUCTION,
    SOURCE_ITEM_LIST,
    SOURCE_NEW,
    SOURCE_OTHER_PLANT,
    SOURCE_OTHER_WBS,
    SOURCE_SAME_PLANT,
    Parameters,
)
from app.ingest.parse import parse_file
from app.jobs.analysis import result_hash
from app.reporting.documents import Templates
from app.reporting.guard import check_report, docx_texts, html_texts, pptx_texts, xlsx_texts

PRG, OWN = "RMS-PRG-FND-CP", "RMS-OWN-ESE-RM"

# Fake previous RBC report ("CPDS Report" sheet): invented plants, items and values.
PREVIOUS = pd.DataFrame(
    [
        ["1000", "101", PRG, "MALARIA", 100, 10000, "Net (fake)", "EA"],
        ["1000", "102", PRG, "CPDS", 50, 5000, "Kit (fake)", "EA"],
        ["7000", "101", PRG, "MALARIA", 20, 2000, "Net (fake)", "EA"],
        ["7000", "103", PRG, "ARV&OI", 10, 1000, "Tablets (fake)", "BOX"],
        ["1000", "104", OWN, "Oncology", 5, 500, "Vial (fake)", "EA"],  # not RBC stock
    ],
    columns=[
        "Valuation Area",
        "Material",
        "WBS",
        "Category",
        "Closing Stock",
        "Closing Value",
        "Material Description",
        "Base Unit of Measure",
    ],
)

GL = pd.DataFrame(
    [
        ["1000", "000101", PRG, 120, 15000],
        ["1000", "101", OWN, 50, 5000],  # not the RBC WBS: left out
        ["1000", "102", PRG, 50, 5000],
        ["1000", "104", PRG, 8, 800],
        ["7000", "102", PRG, 5, 600],
        ["1000", "105", PRG, 3, 300],
        ["1000", "106", PRG, 1, 100],
        ["7000", "107", PRG, 0, 0],  # all zero: dropped
    ],
    columns=["Plant", "Material", "WBS Element", "Quantity", "Amount in Local Currency"],
)
# Fake MB51 of the period (31 March to 30 June 2026).
SCRAP = pd.DataFrame(
    [
        ["1000", "101", "551", PRG, -10, -1000, "15.06.2026"],  # scrapping
        ["1000", "101", "101", PRG, 30, 3000, "10.04.2026"],  # receipt
        ["1000", "101", "102", PRG, -5, -500, "11.04.2026"],  # receipt reversed
        ["1000", "102", "601", PRG, -10, -1000, "20.05.2026"],  # goods issue
        ["7000", "101", "641", PRG, -20, -2000, "02.04.2026"],  # goods issue (STO)
        ["1000", "101", "101", OWN, 50, 5000, "10.04.2026"],  # other WBS: left out
        ["1000", "102", "311", PRG, -1, -100, "10.04.2026"],  # other type: other movements
        ["1000", "102", "601", PRG, -3, -300, "03.07.2026"],  # after the period: left out
    ],
    columns=[
        "Plant",
        "Material",
        "Movement Type",
        "WBS Element",
        "Quantity",
        "Amt.in Loc.Cur.",
        "Posting Date",
    ],
)
ITEMS = pd.DataFrame({"SAP MATERIAL": ["000105"], "Category": ["TB"]})
PLANTS = pd.DataFrame({"Plant": ["1000", "7000"], "Name 1": ["Head office", "Hub"]})

PARAMS = Parameters(opening_date=date(2026, 3, 31), closing_date=date(2026, 6, 30))


def inputs(previous: pd.DataFrame = PREVIOUS) -> dict[str, pd.DataFrame]:
    gl = GL.assign(**{"Material Description": None})
    scrap = SCRAP.assign(**{"Material Description": "Net (fake)"})
    previous = previous.copy()
    for column in ("Material Description", "Base Unit of Measure"):
        if column not in previous:
            previous[column] = None
    return {
        "previous_report": previous,
        "gl_lines": gl,
        "scrapping": scrap,
        "plants": PLANTS,
        "item_categories": ITEMS,
    }


@pytest.fixture(scope="module")
def result() -> Any:
    return ANALYSIS.compute(inputs(), PARAMS)


def line(result: Any, plant: str, material: str) -> dict[str, Any]:
    lines = result.tables["lines"]
    row = lines[(lines["Plant"] == plant) & (lines["Material"] == material)]
    assert len(row) == 1, (plant, material)
    return dict(row.iloc[0])


def test_lines_hand_checked(result: Any) -> None:
    net = line(result, "1000", "101")  # 10,000 -> 15,000 - 1,000 = 14,000
    assert (net["Opening Value"], net["Closing Value"], net["Difference"]) == (10000, 14000, 4000)
    assert net["Difference %"] == pytest.approx(0.4)
    assert (net["Increase / Reduction"], net["Category"]) == (INCREASE, "MALARIA AND NTD")
    assert net["Category Source"] == SOURCE_SAME_PLANT

    kit = line(result, "1000", "102")
    assert (kit["Difference"], kit["Increase / Reduction"]) == (0, NO_CHANGE)

    gone = line(result, "7000", "101")  # opening only
    assert (gone["Closing Value"], gone["Difference %"], gone["Increase / Reduction"]) == (
        0,
        -1.0,
        REDUCTION,
    )

    assert line(result, "7000", "103")["Category"] == "ARV and Opportunistic Infection"

    vial = line(result, "1000", "104")  # on RMS-OWN-ESE-RM before: not in the opening
    assert (vial["Category"], vial["Category Source"]) == ("Cross Cutting", SOURCE_OTHER_WBS)
    assert (vial["Opening Value"], vial["Closing Value"]) == (0, 800)
    assert pd.isna(vial["Difference %"])
    assert line(result, "7000", "102")["Category Source"] == SOURCE_OTHER_PLANT
    listed = line(result, "1000", "105")
    assert (listed["Category"], listed["Category Source"]) == ("TB", SOURCE_ITEM_LIST)
    new = line(result, "1000", "106")
    assert (new["Category"], new["Category Source"]) == ("CPDS", SOURCE_NEW)

    assert len(result.tables["lines"]) == 8  # 7000/107 (all zero) is dropped
    assert net["Plant Name"] == "Head office"


def test_summary_by_programme(result: Any) -> None:
    summary = result.tables["summary"].set_index("Category")
    assert list(summary.index) == [
        "CPDS",
        "MALARIA AND NTD",
        "ARV and Opportunistic Infection",
        "Cross Cutting",
        "TB",
        "Total",
    ]
    assert (summary.at["CPDS", "Opening Value"], summary.at["CPDS", "Closing Value"]) == (
        5000,
        5700,
    )
    assert summary.at["MALARIA AND NTD", "Difference"] == 2000
    assert summary.at["ARV and Opportunistic Infection", "Increase / Reduction"] == REDUCTION
    assert summary.at["Cross Cutting", "Opening Value"] == 0
    assert summary.at["Total", "Opening Value"] == 18000
    assert summary.at["Total", "Closing Value"] == 20800
    assert summary.at["Total", "Difference %"] == pytest.approx(2800 / 18000)
    metrics = result.tables["metrics"].set_index("Category")
    assert metrics.at["CPDS", "Lines in stock"] == 3
    assert metrics.at["Total", "Lines in stock"] == 6
    assert metrics.at["Total", "Share of closing"] == pytest.approx(1.0)
    assert result.extras["position"] == "Increase of 15.6%"
    assert all(c.passed for c in result.checks), [c.label for c in result.checks if not c.passed]
    kpi = {k.label: k.value for k in result.metrics}
    assert kpi["New items to review"] == 1


def test_products_and_closing_lists(result: Any) -> None:
    top = result.tables["top_products"]
    assert list(top["Material"])[:2] == ["101", "102"]  # 14,000 then 5,600
    new = dict(zip(top["Material"], top["New product"], strict=True))
    assert (new["105"], new["104"]) == ("Yes", "Yes")
    assert list(result.tables["top_reductions"]["Material"]) == ["103"]  # 101 nets to +2,000
    closing = result.tables["closing_list"]
    assert list(closing["#"]) == list(range(1, len(closing) + 1))
    assert closing["Closing Value"].sum() == 20800
    by_material = result.tables["by_material"]
    assert by_material.loc[by_material["SAP Material"] == "102", "Plants Holding"].item() == 2


def test_category_names_are_normalised() -> None:
    assert normalise_category("ARV AND OPPORTUNITIC INFECTION") == "ARV and Opportunistic Infection"
    assert normalise_category(" malaria ") == "MALARIA AND NTD"
    assert normalise_category("Program Lab") == "CPDS"
    assert normalise_category("specialized") == "Cross Cutting"
    assert normalise_category("cpds") == "CPDS"
    assert normalise_category("Something new") == "Something new"


def test_same_inputs_same_result() -> None:
    assert result_hash(ANALYSIS.compute(inputs(), PARAMS)) == result_hash(
        ANALYSIS.compute(inputs(), PARAMS)
    )


def test_dates_must_be_in_order() -> None:
    with pytest.raises(ValueError, match="closing date must be after"):
        Parameters(opening_date=date(2026, 6, 30), closing_date=date(2026, 3, 31))


def test_receipts_and_goods_issues_hand_checked(result: Any) -> None:
    """MB51 on the WBS in the period, by movement type (Ferdinand, 26 October 2026)."""
    net = line(result, "1000", "101")  # 101 +3,000, 102 -500; scrapped -1,000
    assert (net["Receipts Qty"], net["Receipts Value"]) == (25, 2500)
    assert (net["Goods Issues Value"], net["Scrapped Value"]) == (0, -1000)
    assert net["Other Movements Value"] == 14000 - 10000 - 2500 + 1000
    kit = line(result, "1000", "102")  # 601 in the period; 311 and the July 601 are not issues
    assert (kit["Goods Issues Qty"], kit["Goods Issues Value"]) == (-10, -1000)
    assert kit["Other Movements Value"] == 1000
    gone = line(result, "7000", "101")  # 641: the whole opening went out
    assert (gone["Goods Issues Value"], gone["Other Movements Value"]) == (-2000, 0)

    summary = result.tables["summary"].set_index("Category")
    total = summary.loc["Total"]
    assert (total["Receipts Value"], total["Goods Issues Value"]) == (2500, -3000)
    assert (total["Scrapped Value"], total["Other Movements Value"]) == (-1000, 4300)
    kpi = {k.label: k.value for k in result.metrics}
    assert (kpi["Receipts"], kpi["Goods issues"], kpi["Scrapped (expired)"]) == (
        2500,
        -3000,
        -1000,
    )
    assert kpi["MB51 lines outside the period (left out)"] == 1
    checks = {c.key: c for c in result.checks}
    assert checks["movements"].expected == -500 and checks["movements"].passed
    assert checks["roll_forward"].passed


def test_reports_show_receipts_and_goods_issues(result: Any) -> None:
    files = ANALYSIS.reports(result, PARAMS, Templates())
    readers = {"xlsx": xlsx_texts, "docx": docx_texts, "pptx": pptx_texts, "html": html_texts}
    for report in files:
        check_report(report.kind, report.content)  # banned wording
        text = " ".join(readers[report.kind](report.content)).lower()
        if report.kind != "pptx":
            for word in ("receipts", "goods issues", "scrapped"):
                assert word in text, (report.filename, word)


def test_report_data_sheet_is_next_quarters_opening(result: Any, tmp_path: Path) -> None:
    """The "CPDS Report" sheet of this report is read back as the next previous report."""
    workbook = ANALYSIS.reports(result, PARAMS, Templates())[0].content
    assert "CPDS Report" in load_workbook(io.BytesIO(workbook), read_only=True).sheetnames
    path = tmp_path / "rbc.xlsx"
    path.write_bytes(workbook)
    sheet = next(s for s in parse_file(path, tmp_path / "parsed") if s.name == "CPDS Report")
    spec = datasets.get("rbc_previous_report")
    matches = [m.dataset for m in datasets.match_sheet(sheet.columns, sheet.name)]
    assert "rbc_previous_report" in matches
    raw = pq.read_table(tmp_path / "parsed" / sheet.parquet_path).to_pandas()
    coerced = datasets.coerce(spec, raw, sheet.header_row)
    assert coerced.ok, [i.message for i in coerced.issues]

    following = Parameters(opening_date=date(2026, 6, 30), closing_date=date(2026, 9, 30))
    next_quarter = ANALYSIS.compute(inputs(previous=coerced.frame), following)
    kpi = {k.label: k.value for k in next_quarter.metrics}
    assert kpi["Opening stock (30 Jun 2026)"] == 20800
    assert kpi["New items to review"] == 0  # every item now has a category
    assert all(c.passed for c in next_quarter.checks)


def test_dashboard_follows_the_rbc_programme_stock_layout(result: Any) -> None:
    """Value bands, products in stock and the closing stock list (Ferdinand's layout)."""
    from app.analyses.rbc_accountability import dashboard_html

    assert dashboard_html.value_band(0, 1_000_000) == "RWF 1M to 10M"
    assert dashboard_html.value_band(-2_000_000_000, 5) == "RWF 1B and over"
    assert dashboard_html.value_band(999_999, 0) == "Under RWF 1M"

    lines = result.tables["lines"]
    closing = dashboard_html.closing_facts(lines)
    held = lines[(lines["Closing Value"].round(2) != 0) | (lines["Closing Qty"].round(3) != 0)]
    assert len(closing) == len(held)
    assert closing["Closing Value"].sum() == pytest.approx(lines["Closing Value"].sum())

    [page] = [r for r in ANALYSIS.reports(result, PARAMS, Templates()) if r.kind == "html"]
    text = " ".join(html_texts(page.content))
    for heading in ("Closing stock list", "Products by plant", "Products in stock"):
        assert heading in text, heading
