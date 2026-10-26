"""Closing stock by WBS: hand-checked figures on fake data (see the table in each test)."""

import io
from datetime import date
from typing import Any

import pandas as pd
import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook

from app.analyses.closing_stock_wbs import (
    ANALYSIS,
    FLAG_EXPIRED,
    FLAG_NEGATIVE,
    FLAG_WBS,
    OTHER,
    OWN,
    PROGRAMS,
    Parameters,
)
from app.analyses.common.closing import NO_WBS
from app.jobs.analysis import result_hash
from app.reporting.documents import Templates
from app.reporting.guard import check_report

GL_HEADERS = ["Plant", "Material", "WBS Element", "Quantity", "Amount in Local Currency"]
MB51_HEADERS = ["Plant", "Material", "Movement Type", "WBS Element", "Quantity", "Amt.in Loc.Cur."]

# Fake G/L line items of the stock accounts (invented plants, materials and values).
GL_ROWS = [
    ["1000", "0000000101", PROGRAMS, 100, 10000],
    ["1000", "101", PROGRAMS, 20, 2000],  # same item: leading zeros do not matter
    ["1000", "101", OWN, 10, 1000],
    ["1000", "102", "", 5, 500],  # no WBS
    ["7000", "201", PROGRAMS, 50, 5000],
    [" 7000", "", "", 0, 16500],  # subtotal line without a material: left out
    ["7000", "202", "ZZ-OTHER-WBS", 4, 400],
]
# Fake scrapping lines (MB51): negative quantities and amounts.
SCRAP_ROWS = [
    ["1000", "101", "551", PROGRAMS, -20, -2000],
    ["7000", "201", "551", OWN, -10, -1000],  # scrapped on a WBS with no stock of the item
    ["7000", "203", "551", PROGRAMS, -3, -300],  # item not in the closing stock
    ["1000", "102", "551", "", -5, -500],
]


def frame(headers: list[str], rows: list[list[Any]], extra: dict[str, Any]) -> pd.DataFrame:
    data = pd.DataFrame(rows, columns=headers)
    for column, value in extra.items():
        data[column] = value
    return data


def inputs() -> dict[str, pd.DataFrame]:
    gl = frame(GL_HEADERS, GL_ROWS, {"Material Description": None, "G/L Account": "11000100"})
    gl.loc[gl["Material"].str.strip("0") == "101", "Material Description"] = "Gloves (fake)"
    scrap = frame(MB51_HEADERS, SCRAP_ROWS, {"Material Description": "Expired item (fake)"})
    plants = pd.DataFrame({"Plant": ["1000", "7000"], "Name 1": ["Head office", "Hub"]})
    return {"gl_lines": gl, "scrapping": scrap, "plants": plants}


PARAMS = Parameters(period_end=date(2026, 6, 30))


@pytest.fixture(scope="module")
def result() -> Any:
    return ANALYSIS.compute(inputs(), PARAMS)


def item(result: Any, plant: str, material: str) -> dict[str, Any]:
    items = result.tables["items"]
    row = items[(items["Plant"] == plant) & (items["Material"] == material)]
    assert len(row) == 1, (plant, material)
    return dict(row.iloc[0])


def test_items_hand_checked(result: Any) -> None:
    # 1000/101: closing 130 / 13,000 (PRG 120 / 12,000, OWN 10 / 1,000), scrapped -20 / -2,000
    gloves = item(result, "1000", "101")
    assert (gloves["Closing Qty"], gloves["Closing Value"]) == (130, 13000)
    assert (gloves["Final Qty"], gloves["Final Value"]) == (110, 11000)
    assert (gloves[OWN], gloves[PROGRAMS], gloves[OTHER]) == (1000, 10000, 0)
    assert gloves["Main WBS"] == PROGRAMS
    assert gloves["All WBS with balance"] == f"{OWN}; {PROGRAMS}"
    assert gloves["Flag"] == FLAG_EXPIRED
    assert gloves["Plant Name"] == "Head office"
    assert gloves["Material Description"] == "Gloves (fake)"
    assert gloves["ID"] == "1000101"

    # 1000/102: all stock scrapped on "no WBS"; the main WBS falls back to the closing value.
    empty = item(result, "1000", "102")
    assert (empty["Final Qty"], empty["Final Value"]) == (0, 0)
    assert empty["Main WBS"] == NO_WBS
    assert empty["Flag"] == FLAG_EXPIRED

    # 7000/201: scrapped on RMS-OWN-ESE-RM, which holds none of it.
    moved = item(result, "7000", "201")
    assert (moved[PROGRAMS], moved[OWN], moved["Final Value"]) == (5000, -1000, 4000)
    assert moved["Flag"] == FLAG_WBS

    # 7000/202: other WBS goes to "No WBS / other".
    other = item(result, "7000", "202")
    assert (other[OTHER], other["Flag"]) == (400, "")

    # 7000/203: scrapped only, not in the closing stock: negative.
    missing = item(result, "7000", "203")
    assert (missing["Final Qty"], missing["Final Value"]) == (-3, -300)
    assert missing["In closing stock"] == "No"
    assert missing["Flag"] == FLAG_NEGATIVE


def test_totals_and_checks(result: Any) -> None:
    kpi = {k.label: k.value for k in result.metrics}
    assert kpi["Closing stock (G/L)"] == 18900
    assert kpi["Expired stock scrapped"] == -3800
    assert kpi["Final closing stock"] == 15100
    assert kpi["Items"] == 5
    assert kpi["Items to check"] == 2
    split = dict(
        zip(result.tables["by_split"]["WBS"], result.tables["by_split"]["Final Value"], strict=True)
    )
    assert split == {OWN: 0, PROGRAMS: 14700, OTHER: 400}
    assert all(c.passed for c in result.checks), [c.label for c in result.checks if not c.passed]
    counts = dict(
        zip(result.tables["counts"]["Count"], result.tables["counts"]["Number"], strict=True)
    )
    assert counts["G/L lines without a material (left out)"] == 1
    assert counts["Items with more than one WBS"] == 2
    assert any("ZZ-OTHER-WBS" in note for note in result.dashboard.notes)
    assert any("16,500" in note for note in result.dashboard.notes)


def test_scrapping_lines_show_whether_the_item_is_in_stock(result: Any) -> None:
    scrap = result.tables["scrapping"]
    in_stock = dict(zip(scrap["Material"], scrap["In closing stock?"], strict=True))
    assert in_stock == {"101": "Yes", "201": "Yes", "203": "No", "102": "Yes"}


def test_same_inputs_same_result() -> None:
    first = ANALYSIS.compute(inputs(), PARAMS)
    second = ANALYSIS.compute(inputs(), PARAMS)
    assert result_hash(first) == result_hash(second)


def test_plant_names_are_optional() -> None:
    data = inputs()
    del data["plants"]
    result = ANALYSIS.compute(data, PARAMS)
    assert set(result.tables["items"]["Plant Name"]) == {""}
    assert all(c.passed for c in result.checks)


def test_reports_pass_the_language_guard(result: Any) -> None:
    files = ANALYSIS.reports(result, PARAMS, Templates())
    assert [f.kind for f in files] == ["xlsx", "html", "docx", "pptx"]
    for report in files:
        check_report(report.kind, report.content)
    workbook = load_workbook(io.BytesIO(files[0].content), read_only=True)
    assert "Final Closing Stock" in workbook.sheetnames
    assert "Expired Scrapped" in workbook.sheetnames


@pytest.mark.usefixtures("jobs")
def test_runs_end_to_end_on_uploaded_csv_files(logged_in: TestClient) -> None:
    from tests.test_engine import upload

    def csv(headers: list[str], rows: list[list[Any]]) -> bytes:
        lines = [",".join(headers)] + [",".join(str(v) for v in row) for row in rows]
        return ("\n".join(lines) + "\n").encode()

    gl = upload(logged_in, csv(GL_HEADERS, GL_ROWS), "gl_stock.csv")
    scrap = upload(logged_in, csv(MB51_HEADERS, SCRAP_ROWS), "mb51_scrapping.csv")
    assert gl["sheets"][0]["matches"][0]["dataset"] == "gl_stock_lines"
    assert scrap["sheets"][0]["matches"][0]["dataset"] == "mb51_material_documents"
    response = logged_in.post(
        "/api/admin/runs",
        json={
            "analysis": "closing_stock_wbs",
            "label": "June closing",
            "parameters": {"period_end": "2026-06-30"},
            "inputs": {
                "gl_lines": {"data_file_id": gl["id"], "sheet": "CSV"},
                "scrapping": {"data_file_id": scrap["id"], "sheet": "CSV"},
            },
        },
    )
    assert response.status_code == 201, response.text
    run = logged_in.get(f"/api/admin/runs/{response.json()['id']}").json()
    assert run["status"] == "succeeded", run["errors"] or run["message"]
    assert run["message"] == "All checks passed."
    final = next(k for k in run["metrics"] if k["label"] == "Final closing stock")
    assert final["value"] == 15100


def test_only_scrapping_movement_types_are_deducted() -> None:
    """MB51 holds every movement; only 551 and 555 count as scrapping by default."""
    data = inputs()
    others = pd.DataFrame(
        [
            ["1000", "101", "601", PROGRAMS, -50, -5000],  # delivery: not scrapping
            ["1000", "101", "555", PROGRAMS, -1, -100],  # scrapping from blocked stock
        ],
        columns=MB51_HEADERS,
    ).assign(**{"Material Description": "Expired item (fake)"})
    data["scrapping"] = pd.concat([data["scrapping"], others], ignore_index=True)
    result = ANALYSIS.compute(data, PARAMS)
    kpi = {k.label: k.value for k in result.metrics}
    assert kpi["Expired stock scrapped"] == -3800 - 100
    counts = dict(
        zip(result.tables["counts"]["Count"], result.tables["counts"]["Number"], strict=True)
    )
    assert counts["MB51 lines of other movement types (left out)"] == 1
    everything = ANALYSIS.compute(
        data, Parameters(period_end=date(2026, 6, 30), scrapping_movement_types="551, 555, 601")
    )
    assert {k.label: k.value for k in everything.metrics}["Expired stock scrapped"] == -8900


def test_counted_stock_without_expired() -> None:
    """Legacy w5 rules 8-9: scrapped stock comes off the counted stock when it covers it."""
    data = inputs()
    data["counted"] = pd.DataFrame(
        {
            "ValA": ["1000", "7000", "7000"],
            "Material": ["101", "201", "203"],
            "Counted Quantities": [130, 50, 2],  # 203: 2 counted < 3 scrapped: unchanged
            "Counted Value": [13000, 5000, 200],
            "Reasons for differences where Applicable": ["NO VARIANCE", "NO VARIANCE", "x"],
        }
    )
    result = ANALYSIS.compute(data, PARAMS)
    gloves = item(result, "1000", "101")  # scrapped -20 / -2,000 comes off
    assert (gloves["Counted Qty excl. Expired"], gloves["Counted Value excl. Expired"]) == (
        110,
        11000,
    )
    assert gloves["Difference after Scrapping"] == 0
    moved = item(result, "7000", "201")  # scrapped -10 / -1,000
    assert moved["Counted Value excl. Expired"] == 4000
    short = item(result, "7000", "203")
    assert (short["Counted Qty excl. Expired"], short["Difference after Scrapping"]) == (2, 500)
    missing = item(result, "7000", "202")  # not in the count
    assert (missing["Counted Value"], missing["Difference after Scrapping"]) == (0, -400)
    kpi = {k.label: k.value for k in result.metrics}
    assert kpi["Counted value excl. expired"] == 11000 + 4000 + 200
    assert all(c.passed for c in result.checks)
