import io
from pathlib import Path

import pandas as pd
import pytest
from docx import Document

from app.reporting.documents import (
    ChartBlock,
    Highlight,
    KpiCard,
    Recommendation,
    ReportDocument,
    Section,
    TableBlock,
    format_value,
)
from app.reporting.guard import (
    BannedWordingError,
    check_report,
    docx_texts,
    find_banned,
    pptx_texts,
)
from app.reporting.slides import build_pptx
from app.reporting.word import TemplateError, build_docx

DOC = ReportDocument(
    title="Stock report",
    subtitle="Fake data",
    date_text="30 June 2026",
    sections=[
        Section(
            "Summary",
            paragraphs=["Total stock is RWF 1,000."],
            tables=[
                TableBlock(
                    "By plant",
                    pd.DataFrame({"Plant": ["P1", "P2"], "Value": [600.0, 400.0]}),
                    {"Value": "money"},
                    total_row=True,
                )
            ],
        )
    ],
)


def test_banned_wording_is_found() -> None:
    assert find_banned(["Prepared with AI support"])
    assert find_banned(["Amount To Be Confirmed later"])
    assert find_banned(["AID kit", "Kigali", "maintain", "the ai word in lower case"]) == []


def test_values_are_formatted_for_reports() -> None:
    assert format_value(-1234.4, "money") == "(1,234)"
    assert format_value(0.0891, "percent") == "8.9%"
    assert format_value(None, "money") == ""


# The parts of a management (SMT) deck.
DECK = ReportDocument(
    title="Stock report",
    subtitle="Fake data",
    date_text="30 June 2026",
    sections=[
        Section(
            "By plant",
            message="Plant P1 holds most of the stock",
            takeaway="Count P1 first.",
            charts=[
                ChartBlock(
                    "Value by plant", ["P1", "P2"], {"Value": [600.0, 400.0]}, "bar", "money"
                )
            ],
            tables=[
                TableBlock(
                    "Plants",
                    pd.DataFrame({"Plant": ["P1", "P2"], "Value": [600.0, 400.0]}),
                    {"Value": "money"},
                )
            ],
        )
    ],
    summary_message="RWF 1,000 of stock in two plants",
    kpis=[KpiCard("Stock value", "RWF 1,000", "2 plants"), KpiCard("Plants", "2", tone="green")],
    highlights=[Highlight("1 · Concentration", "P1 holds 60%.", "yellow")],
    recommendations=[Recommendation("Count plant P1", "Inventory control", "7 days")],
    notes=[Highlight("Coverage", "Fake files only.")],
    audience="Senior Management Team",
    draft=True,
)


def test_word_report_without_template() -> None:
    content = build_docx(DOC)
    texts = docx_texts(content)
    assert "STOCK REPORT" in " | ".join(texts)  # the cover
    assert "Total" in texts and "1,000" in texts  # total row of the table
    check_report("docx", content)


def test_word_report_with_the_deck_parts() -> None:
    content = build_docx(DECK)
    text = " | ".join(docx_texts(content))
    for part in (
        "Executive summary",
        "RWF 1,000 of stock in two plants",
        "Stock value",
        "1 · Concentration",
        "Plant P1 holds most of the stock",
        "So what",
        "Count P1 first.",
        "Recommendations and next steps",
        "Inventory control",
        "Data notes",
        "DRAFT",
    ):
        assert part in text, part
    check_report("docx", content)


def _template(tmp_path: Path, body: str) -> Path:
    document = Document()
    document.add_paragraph("RMS LETTERHEAD")
    document.add_paragraph("{{ title }} - {{ report_date }}")
    document.add_paragraph(body)
    document.add_paragraph("End of template")
    path = tmp_path / "template.docx"
    document.save(str(path))
    return path


def test_word_report_on_an_organisation_template(tmp_path: Path) -> None:
    content = build_docx(DOC, _template(tmp_path, "{{p body }}"))
    texts = docx_texts(content)
    assert texts[0] == "RMS LETTERHEAD"
    assert "Stock report - 30 June 2026" in texts
    assert "Total stock is RWF 1,000." in texts
    assert texts[-1] == "End of template" or "End of template" in texts


def test_word_template_without_body_placeholder_is_refused(tmp_path: Path) -> None:
    with pytest.raises(TemplateError, match="body"):
        build_docx(DOC, _template(tmp_path, "no placeholder here"))


def test_powerpoint_deck() -> None:
    content = build_pptx(DOC)
    texts = pptx_texts(content)
    assert "STOCK REPORT" in texts  # the cover
    assert "By plant" in texts
    assert "RMS · Stock report · 30 June 2026" in texts  # the footer of every slide
    check_report("pptx", content)


def test_powerpoint_management_deck() -> None:
    from pptx import Presentation

    content = build_pptx(DECK)
    texts = " | ".join(pptx_texts(content))
    for part in (
        "Executive summary",
        "RWF 1,000",
        "So what",
        "Count P1 first.",
        "Recommendations and next steps",
        "Count plant P1",
        "Data notes",
        "Thank You",
        "Draft for SMT",
    ):
        assert part in texts, part
    deck = Presentation(io.BytesIO(content))
    charts = [s for slide in deck.slides for s in slide.shapes if s.has_chart]
    assert len(charts) == 1  # a native, editable chart
    check_report("pptx", content)


def test_guard_rejects_a_report_with_banned_wording() -> None:
    bad = ReportDocument("Plan", "AI generated summary", "30 June 2026", sections=[])
    with pytest.raises(BannedWordingError, match="AI"):
        check_report("docx", build_docx(bad))
    buffer = io.BytesIO(build_pptx(bad))
    with pytest.raises(BannedWordingError):
        check_report("pptx", buffer.getvalue())


def test_long_dashboard_tables_are_capped() -> None:
    from app.analyses.dashboard import MAX_DASHBOARD_ROWS, Column, Dashboard, Table
    from app.jobs.analysis import cap_dashboard

    rows = [{"n": i} for i in range(MAX_DASHBOARD_ROWS + 10)]
    dashboard = Dashboard(
        title="t",
        tables=[Table(id="x", title="x", columns=[Column(key="n", label="n")], rows=rows)],
    )
    capped = cap_dashboard(dashboard).tables[0]
    assert len(capped.rows) == MAX_DASHBOARD_ROWS
    assert "All lines are in the Excel workbook" in capped.note
    assert len(dashboard.tables[0].rows) == MAX_DASHBOARD_ROWS + 10  # original untouched


def test_excel_sheet_row_limit(monkeypatch: pytest.MonkeyPatch) -> None:
    import io as _io

    from openpyxl import load_workbook

    from app.reporting import excel

    monkeypatch.setattr(excel, "EXCEL_MAX_ROWS", 20)
    frame = pd.DataFrame({"n": range(50)})
    content = excel.build_workbook("t", "s", [], [TableBlock("Lines", frame)], [])
    sheet = load_workbook(_io.BytesIO(content))["Lines"]
    assert "showing the first 15 of 50" in str(sheet["A3"].value)
    assert sheet.max_row <= 20
