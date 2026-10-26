"""Spot-check variance follow-up: bin list, cases, secure link, e-mails (fake data)."""

import io
from datetime import UTC, date, datetime, timedelta
from typing import Any

import pandas as pd
import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook
from sqlalchemy import select
from sqlalchemy.orm import Session

from app import variances
from app.core.config import Settings
from app.db.models import BinOwner, VarianceCase, VarianceStaff
from app.jobs import variances as variance_jobs
from tests.conftest import FakeMailer
from tests.test_engine import upload

# Fake bin list: names written in different orders and spellings, two unit bins,
# the e-mail block below the bins with its heading.
BIN_LIST = pd.DataFrame(
    {
        "RESPONSIBLE PERSON": [
            "ALICE MUKAMANA",
            "Mukamana Alice",
            "MUKAMANA ALYCE",
            "BOB KALISA",
            "RECEIVING",
            "PICKING",
            "Name",
            "Alice Mukamana",
            "Kalisa Bob",
            "a@example.com b@example.com",
        ],
        "Storage Bin": [
            "A-01",
            "a-02",
            "A-03",
            "B-01",
            "R-00",
            "P-00",
            "email",
            "alice@example.com",
            "bob@example.com",
            None,
        ],
        "Storage Type": ["7000", "7000", "7000", "7010", "7050", "7060", None, None, None, None],
        "Storage Section": ["S1", "S1", "S1", "S2", None, None, None, None, None, None],
    }
)


def settings() -> Settings:
    return Settings(
        _env_file=None,
        smtp_host="smtp.example.com",
        mail_from="stock@example.com",
        variance_bcc="manager@example.com",
        site_url="https://rms.example.com",
    )


def test_names_match_whatever_the_order_and_small_spelling_differences() -> None:
    assert variances.person_key("Mukamana Alice") == variances.person_key("ALICE  MUKAMANA")
    assert variances.same_person("ALICE MUKAMANA", "MUKAMANA ALYCE")
    assert not variances.same_person("ALICE MUKAMANA", "BOB KALISA")


def test_bin_list_import(db: Session) -> None:
    result = variances.import_bin_list(db, BIN_LIST, settings())
    db.commit()
    assert sorted(result.staff) == ["Alice Mukamana", "Bob Kalisa"]
    assert (result.bins, result.unit_bins, result.emails) == (4, 2, 2)
    owners = {b.storage_bin: b.staff_id for b in db.scalars(select(BinOwner))}
    staff = {s.name: s for s in db.scalars(select(VarianceStaff))}
    alice = staff["Alice Mukamana"]
    assert alice.email == "alice@example.com"
    assert owners == {
        "A-01": alice.id,
        "A-02": alice.id,
        "A-03": alice.id,
        "B-01": staff["Bob Kalisa"].id,
    }
    # Importing a new list keeps the same people (no duplicates).
    variances.import_bin_list(db, BIN_LIST, settings())
    db.commit()
    assert len(db.scalars(select(VarianceStaff)).all()) == 2


SPOT = pd.DataFrame(
    {
        "Storage Bin": ["A-01", "B-01", "A-02", "Z-99"],
        "Product": ["1001", "1002", "1003", "1004"],
        "Variance": [-2.0, 0.0, None, 5.0],
        "Product Description": ["Gloves", "Masks", "Kits", "Tabs"],
        "Batch": ["B1", None, None, None],
        "Avail. Qty in AUoM": [10.0, 5.0, 8.0, 1.0],
        "Counted Quantities": [8.0, 5.0, 9.0, 6.0],
        "Stock Reference Doc.": [None] * 4,
    }
)
PICKING = pd.DataFrame(
    {
        "BIN LOCATION": ["B-01", "B-01"],
        "PRODUCT CODE": [1002.0, 1002.0],
        "Quantity": [3.0, 3.0],  # the same line twice in the register
        "PRODUCT DESCRIPTION": ["Masks", "Masks"],
        "BATCH": [None, None],
        "Date reported": pd.to_datetime(["2026-01-12", "2026-01-12"]),
        "REPORTER/PICKING STAFF": ["CHARLES", "CHARLES"],
    }
)
PI = pd.DataFrame(
    {
        "Phys. Inventory Doc.": ["700001", "700001"],
        "Storage Bin": ["A-03", "A-01"],
        "Book Quantity": [10.0, 4.0],
        "Counted Quantity": [7.0, 4.0],
        "Difference Quantity": [-3.0, 0.0],
        "Product": ["1005", "1001"],
        "Item": ["1", "2"],
        "Product Description": ["Syringes", "Gloves"],
        "Batch": [None, None],
        "Count Date": pd.to_datetime(["2026-01-10", "2026-01-10"]),
        "Counter": ["COUNTER1", "COUNTER1"],
        "Base Unit of Measure": ["EA", "EA"],
    }
)


def test_cases_from_each_source(db: Session) -> None:
    variances.import_bin_list(db, BIN_LIST, settings())
    spot = variances.import_cases(db, variances.SPOT_CHECK, SPOT, None, "sha-1")
    # A-01 (-2), A-02 (counted 9 - available 8 = +1), Z-99 (+5, no owner); B-01 has no variance
    assert (spot.created, spot.no_variance, spot.unassigned) == (3, 1, 1)
    picking = variances.import_cases(db, variances.PICKING, PICKING, None, "sha-2")
    assert (picking.created, picking.duplicates) == (1, 1)
    pi = variances.import_cases(db, variances.PHYSICAL_INVENTORY, PI, None, "sha-3")
    assert (pi.created, pi.no_variance) == (1, 1)
    db.commit()
    cases = {(c.source, c.storage_bin): c for c in db.scalars(select(VarianceCase))}
    assert cases[("spot_check", "A-02")].difference_qty == 1
    missing = cases[("picking", "B-01")]
    assert missing.difference_qty == -3
    assert missing.product == "1002"  # not "1002.0"
    assert missing.found_on == date(2026, 1, 12)
    assert missing.reported_by == "CHARLES"
    counted = cases[("physical_inventory", "A-03")]
    assert counted.reference == "700001/1"
    assert (counted.system_qty, counted.counted_qty) == (10, 7)
    assert cases[("spot_check", "Z-99")].staff_id is None
    # The register grows: importing it again adds nothing; another spot check file is new.
    assert variances.import_cases(db, variances.PICKING, PICKING, None, "sha-4").created == 0
    assert variances.import_cases(db, variances.SPOT_CHECK, SPOT, None, "sha-5").created == 3
    assert len({c.token for c in db.scalars(select(VarianceCase))}) == 8  # one token each


def test_colours_and_overdue() -> None:
    s = settings()
    today = date(2026, 1, 15)
    case = VarianceCase(status="open", created_at=datetime(2026, 1, 13, 8, tzinfo=UTC))
    assert variances.colour(case, today, s) == "amber"  # open 2 days
    case.created_at = datetime(2026, 1, 12, 8, tzinfo=UTC)
    assert variances.colour(case, today, s) == "red"  # more than 2 days
    case.status = "corrected"
    assert variances.colour(case, today, s) == "green"


def test_digest_only_on_working_days_from_the_reminder_time() -> None:
    s = settings()
    thursday = datetime(2026, 1, 15, 7, 29)
    assert not variances.digest_due(s, thursday)
    assert variances.digest_due(s, thursday.replace(minute=30))
    assert not variances.digest_due(s, datetime(2026, 1, 17, 9, 0))  # Saturday


def test_correct_needs_a_known_unit(db: Session) -> None:
    case = VarianceCase(status="open")
    with pytest.raises(ValueError, match="Unknown responsible unit"):
        variances.correct(case, "Recounted", "Somebody", "", "", datetime.now(UTC))
    variances.correct(
        case,
        "Picked from the wrong bin",
        "Picking",
        "Moved back",
        "Alice",
        datetime(2026, 1, 15, tzinfo=UTC),
    )
    assert case.status == "corrected"
    with pytest.raises(variances.CaseClosed):
        variances.correct(case, "Again", "Picking", "", "", datetime.now(UTC))


def test_emails_go_to_the_owner_with_the_manager_in_bcc(db: Session) -> None:
    s = settings()
    variances.import_bin_list(db, BIN_LIST, s)
    result = variances.import_cases(db, variances.SPOT_CHECK, SPOT, None, "sha-1")
    db.commit()
    messages = variances.assignment_messages(db, s, result.new_cases, date(2026, 1, 15))
    assert [m["To"] for m in messages] == ["alice@example.com"]  # Z-99 has no owner
    mail = messages[0]
    assert mail["Bcc"] == "manager@example.com"
    body = mail.get_content()
    assert "Bin A-01" in body and "Bin A-02" in body
    assert "https://rms.example.com/variance/" in body

    # Morning list: the open cases, overdue ones marked; once a day per person.
    case = db.scalars(select(VarianceCase).where(VarianceCase.storage_bin == "A-01")).one()
    case.created_at = datetime(2026, 1, 10, tzinfo=UTC)
    db.commit()
    mailer = FakeMailer()
    morning = datetime(2026, 1, 15, 6, 0, tzinfo=UTC)  # 08:00 in Kigali
    assert variance_jobs.send_digests(db, s, morning, mailer) == 1
    digest = mailer.sent[0]
    assert digest["Subject"] == "Open stock variances: 2 (1 overdue)"
    assert digest["Bcc"] == "manager@example.com"
    assert "[OVERDUE]" in digest.get_content()
    assert variance_jobs.send_digests(db, s, morning + timedelta(hours=1), mailer) == 0


# --- Through the API ------------------------------------------------------------------------


def _csv(frame: pd.DataFrame) -> bytes:
    return frame.to_csv(index=False).encode()


@pytest.mark.usefixtures("jobs")
def test_admin_import_link_and_correction(
    logged_in: TestClient, client: TestClient, monkeypatch: pytest.MonkeyPatch, settings: Settings
) -> None:
    settings.variance_bcc = "manager@example.com"
    mailer = FakeMailer()
    monkeypatch.setattr(variance_jobs, "get_mailer", lambda: mailer)
    monkeypatch.setattr(variance_jobs, "get_settings", lambda: settings)

    bin_file = upload(logged_in, _csv(BIN_LIST), "bins.csv")
    response = logged_in.post(
        "/api/admin/variances/bin-list", json={"data_file_id": bin_file["id"]}
    )
    assert response.status_code == 200, response.text
    assert response.json()["bins"] == 4

    spot_file = upload(logged_in, _csv(SPOT.drop(columns="Stock Reference Doc.")), "spot.csv")
    response = logged_in.post("/api/admin/variances/import", json={"data_file_id": spot_file["id"]})
    assert response.status_code == 200, response.text
    body = response.json()
    assert (body["source"], body["created"], body["unassigned"]) == ("spot_check", 3, 1)
    assert [m["To"] for m in mailer.sent] == ["alice@example.com"]  # sent by the worker job
    assert mailer.sent[0]["Bcc"] == "manager@example.com"

    overview = logged_in.get("/api/admin/variances").json()
    assert overview["total"] == 3
    assert {c["colour"] for c in overview["cases"]} == {"amber"}
    case = next(c for c in overview["cases"] if c["storage_bin"] == "A-01")
    link = logged_in.get(f"/api/admin/variances/cases/{case['id']}/link").json()["link"]
    token = link.rsplit("/", 1)[-1]

    # The person opens the link: no login, this case only.
    public = client.get(f"/api/public/variances/{token}")
    assert public.status_code == 200
    assert public.json()["storage_bin"] == "A-01"
    assert "token" not in public.json() and "staff_id" not in public.json()
    correction = {
        "reason": "Picked from the wrong bin",
        "responsible_unit": "Picking",
        "action_taken": "Stock moved back",
        "corrected_by": "Alice",
    }
    done = client.post(f"/api/public/variances/{token}/correct", json=correction)
    assert done.status_code == 200, done.text
    assert done.json()["colour"] == "green"
    assert client.post(f"/api/public/variances/{token}/correct", json=correction).status_code == 409
    assert client.get("/api/public/variances/not-a-real-token-at-all").status_code == 404

    export = logged_in.get("/api/admin/variances/export.xlsx")
    assert export.status_code == 200
    sheet = load_workbook(io.BytesIO(export.content))["Variances"]
    texts = {str(c.value) for row in sheet.iter_rows() for c in row if c.value is not None}
    assert "Picked from the wrong bin" in texts


def test_admin_routes_need_login(client: TestClient) -> None:
    for method, path in [
        ("get", "/api/admin/variances"),
        ("post", "/api/admin/variances/import"),
        ("get", "/api/admin/variances/export.xlsx"),
    ]:
        response: Any = getattr(client, method)(path)
        assert response.status_code == 401
