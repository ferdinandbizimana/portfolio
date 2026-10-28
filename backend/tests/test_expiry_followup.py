"""Monthly expiry follow-up: review, justifications, e-mail and secure link (fake data)."""

import io
from datetime import UTC, date, datetime
from typing import Any

import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook
from sqlalchemy import select
from sqlalchemy.orm import Session

from app import expiry
from app.analyses.expiry_monitoring import ANALYSIS
from app.core.config import Settings
from app.db.models import AnalysisRun, ExpiryJustification
from app.jobs.expiry import send_monthly
from app.reporting.guard import check_report
from app.storage import LocalStorage
from tests.conftest import FakeMailer
from tests.test_expiry_monitoring import PARAMS, inputs


def mail_settings(**extra: Any) -> Settings:
    values: dict[str, Any] = {
        "smtp_host": "smtp.example.com",
        "mail_from": "stock@example.com",
        "site_url": "https://rms.example.com",
        "expiry_mail_to": "pharmacist@example.com",
        "expiry_mail_cc": "manager.one@example.com, manager.two@example.com",
        "expiry_mail_bcc": "copy@example.com",
    }
    return Settings(_env_file=None, **(values | extra))


def stored_run(db: Session, storage: LocalStorage) -> AnalysisRun:
    """A finished expiry monitoring run with its tables, as the worker stores them."""
    result = ANALYSIS.compute(inputs(), PARAMS)
    run = AnalysisRun(
        analysis_key="expiry_monitoring",
        analysis_version=ANALYSIS.version,
        label="Quarter",
        parameters=PARAMS.model_dump(mode="json"),
        inputs={},
        status="succeeded",
    )
    db.add(run)
    db.commit()
    folder = storage.path(f"runs/{run.id}/tables")
    folder.mkdir(parents=True, exist_ok=True)
    for name, frame in result.tables.items():
        frame.to_parquet(folder / f"{name}.parquet", index=False)
    return run


def test_review_lists_each_expired_material_once(db: Session, storage: LocalStorage) -> None:
    run = stored_run(db, storage)
    review = expiry.ensure_review(db, storage, run)
    db.commit()
    assert expiry.ensure_review(db, storage, run).id == review.id  # made once
    items = db.scalars(
        select(ExpiryJustification)
        .where(ExpiryJustification.review_id == review.id)
        .order_by(ExpiryJustification.id)
    ).all()
    # Expired at closing: M6 (500), M2 (400), M3 (50), largest first.
    expired = [(i.material, i.value) for i in items if i.kind == expiry.EXPIRED]
    assert expired == [("M6", 500), ("M2", 400), ("M3", 50)]
    # Expiring within 6 months, for the measures: the soonest first.
    projected = [(i.material, i.band, i.value) for i in items if i.kind == expiry.PROJECTED]
    assert projected == [("M4", expiry.WITHIN_3, 60), ("M5", expiry.THREE_TO_6, 120)]
    assert review.closing_date == date(2026, 6, 30)


def test_justification_rules() -> None:
    item = ExpiryJustification(material="M1", value=1.0)
    now = datetime(2026, 7, 1, tzinfo=UTC)
    with pytest.raises(expiry.JustificationError):
        expiry.justify(item, "Not in the list", "", "", now)
    with pytest.raises(expiry.JustificationError):
        expiry.justify(item, expiry.OTHER_CAUSE, " ", "", now)  # Other needs a reason
    expiry.justify(item, expiry.ROOT_CAUSES[0], "", "Pharmacist", now)
    assert (item.root_cause, item.justified_by) == (expiry.ROOT_CAUSES[0], "Pharmacist")

    # Items expiring soon take a measure, not a root cause.
    soon = ExpiryJustification(material="M4", value=1.0, kind=expiry.PROJECTED)
    with pytest.raises(expiry.JustificationError):
        expiry.justify(soon, expiry.ROOT_CAUSES[0], "", "", now)
    with pytest.raises(expiry.JustificationError):
        expiry.justify(soon, expiry.OTHER_MEASURE, "", "", now)  # Other needs details
    expiry.justify(soon, expiry.MEASURES[0], "Two health centres", "Pharmacist", now)
    assert (soon.root_cause, soon.reason) == (expiry.MEASURES[0], "Two health centres")


def test_due_on_the_last_day_of_the_month_from_the_time() -> None:
    settings = mail_settings()
    assert expiry.mail_due(settings, datetime(2026, 10, 31, 8, 0))
    assert not expiry.mail_due(settings, datetime(2026, 10, 31, 7, 59))
    assert not expiry.mail_due(settings, datetime(2026, 10, 30, 9, 0))
    assert expiry.mail_due(settings, datetime(2028, 2, 29, 9, 0))  # leap year


def test_monthly_mail_once_a_month(db: Session, storage: LocalStorage) -> None:
    stored_run(db, storage)
    settings = mail_settings(timezone="UTC")
    mailer = FakeMailer()
    last_day = datetime(2026, 10, 31, 9, 0, tzinfo=UTC)
    assert not send_monthly(db, settings, datetime(2026, 10, 30, 9, 0, tzinfo=UTC), mailer, storage)
    assert send_monthly(db, settings, last_day, mailer, storage)
    assert not send_monthly(db, settings, last_day, mailer, storage)  # already sent this month
    [mail] = mailer.sent
    assert mail["To"] == "pharmacist@example.com"
    assert mail["Cc"] == "manager.one@example.com, manager.two@example.com"
    assert mail["Bcc"] == "copy@example.com"
    assert mail["Subject"] == "Expiry follow-up 2026-10: expired RWF 950, within 3 months RWF 60"
    body = mail.get_body(("plain",)).get_content()  # type: ignore[union-attr]
    assert "- 1000 Head office: expired RWF 900, within 3 months RWF 60" in body
    assert "https://rms.example.com/expiry/" in body
    [attached] = list(mail.iter_attachments())
    content = attached.get_content()
    check_report("xlsx", content)
    sheets = load_workbook(io.BytesIO(content)).sheetnames
    assert sheets[:1] and len(sheets) >= 4


def test_no_mail_without_a_recipient(db: Session, storage: LocalStorage) -> None:
    stored_run(db, storage)
    settings = mail_settings(expiry_mail_to="", timezone="UTC")
    assert not send_monthly(
        db, settings, datetime(2026, 10, 31, 9, 0, tzinfo=UTC), FakeMailer(), storage
    )


def test_admin_and_secure_link(logged_in: TestClient, db: Session, storage: LocalStorage) -> None:
    run = stored_run(db, storage)
    overview = logged_in.get("/api/admin/expiry").json()
    assert [r["run_id"] for r in overview["runs"]] == [run.id]
    assert overview["runs"][0]["review"] is None
    detail = logged_in.post(f"/api/admin/expiry/runs/{run.id}").json()
    assert (detail["review"]["items"], detail["review"]["projected"]) == (3, 2)
    assert detail["measures"] == expiry.MEASURES
    assert [p["plant"] for p in detail["outlook"]] == ["1000", "7001", "7002"]
    token = detail["review"]["link"].rsplit("/", 1)[-1]

    # The pharmacist, without an account.
    logged_in.cookies.clear()
    seen = logged_in.get(f"/api/public/expiry/{token}")
    assert seen.status_code == 200
    item = seen.json()["items"][0]
    refused = logged_in.post(
        f"/api/public/expiry/{token}/items/{item['id']}",
        json={"root_cause": expiry.OTHER_CAUSE, "reason": ""},
    )
    assert refused.status_code == 422
    saved = logged_in.post(
        f"/api/public/expiry/{token}/items/{item['id']}",
        json={
            "root_cause": expiry.ROOT_CAUSES[2],
            "reason": "Demand fell after the change of protocol",
            "justified_by": "Sales pharmacist",
        },
    )
    assert saved.status_code == 200, saved.text
    assert saved.json()["root_cause"] == expiry.ROOT_CAUSES[2]
    soon = next(i for i in seen.json()["items"] if i["kind"] == "projected")
    measure = logged_in.post(
        f"/api/public/expiry/{token}/items/{soon['id']}",
        json={"root_cause": expiry.MEASURES[1], "reason": "To the branch that uses it"},
    )
    assert measure.status_code == 200, measure.text
    assert logged_in.get("/api/public/expiry/not-a-valid-token-at-all-123").status_code == 404
    assert logged_in.get("/api/admin/expiry").status_code == 401


def test_justification_reports(logged_in: TestClient, db: Session, storage: LocalStorage) -> None:
    run = stored_run(db, storage)
    review_id = logged_in.post(f"/api/admin/expiry/runs/{run.id}").json()["review"]["id"]
    workbook = logged_in.get(f"/api/admin/expiry/reviews/{review_id}/justifications.xlsx")
    assert workbook.status_code == 200
    check_report("xlsx", workbook.content)
    sheets = load_workbook(io.BytesIO(workbook.content)).sheetnames
    assert "Measures to avoid expiry" in sheets
    page = logged_in.get(f"/api/admin/expiry/reviews/{review_id}/justifications.html")
    assert page.status_code == 200
    check_report("html", page.content)
    assert b"<title>Expiry justifications: 30 June 2026</title>" in page.content
