"""The analysis engine end to end through the API, with jobs run immediately."""

import io
from typing import Any

import pytest
from docx import Document
from fastapi.testclient import TestClient
from openpyxl import Workbook, load_workbook

from app.api import uploads
from app.reporting.guard import docx_texts
from tests.samples import DEMO_ROWS, demo_csv, demo_workbook

pytestmark = pytest.mark.usefixtures("jobs")


def upload(client: TestClient, content: bytes, name: str, chunk_size: int = 1000) -> dict[str, Any]:
    """Upload in several chunks, the way the browser does."""
    uploads.CHUNK_SIZE = chunk_size
    start = client.post("/api/admin/uploads", json={"filename": name, "size_bytes": len(content)})
    assert start.status_code == 201, start.text
    state = start.json()
    for index in range(state["chunk_count"]):
        piece = content[index * chunk_size : (index + 1) * chunk_size]
        response = client.put(f"/api/admin/uploads/{state['id']}/chunks/{index}", content=piece)
        assert response.status_code == 204, response.text
    done = client.post(f"/api/admin/uploads/{state['id']}/complete")
    assert done.status_code == 201, done.text
    return client.get(f"/api/admin/data-files/{done.json()['id']}").json()  # type: ignore[no-any-return]


@pytest.fixture(autouse=True)
def _restore_chunk_size() -> Any:
    original = uploads.CHUNK_SIZE
    yield
    uploads.CHUNK_SIZE = original


def start_run(
    client: TestClient, file_id: int, sheet: str = "Stock", **extra: Any
) -> dict[str, Any]:
    body = {
        "analysis": "demo_stock_summary",
        "label": "End of June 2026",
        "parameters": {"period_end": "2026-06-30"},
        "inputs": {"stock": {"data_file_id": file_id, "sheet": sheet}},
        **extra,
    }
    response = client.post("/api/admin/runs", json=body)
    assert response.status_code == 201, response.text
    return client.get(f"/api/admin/runs/{response.json()['id']}").json()  # type: ignore[no-any-return]


def test_upload_in_chunks_is_parsed_and_recognised(logged_in: TestClient) -> None:
    data_file = upload(logged_in, demo_workbook(), "MB52 export.xlsx")

    assert data_file["status"] == "ready", data_file["message"]
    stock = next(s for s in data_file["sheets"] if s["name"] == "Stock")
    assert stock["row_count"] == len(DEMO_ROWS)
    assert stock["matches"] == [
        {"dataset": "demo_stock_balance", "title": "Demo: stock balance (fake data)"}
    ]
    notes = next(s for s in data_file["sheets"] if s["name"] == "Notes")
    assert notes["matches"] == []


def test_upload_can_resume_and_rejects_bad_chunks(logged_in: TestClient) -> None:
    uploads.CHUNK_SIZE = 1000
    content = demo_workbook()
    state = logged_in.post(
        "/api/admin/uploads", json={"filename": "x.xlsx", "size_bytes": len(content)}
    ).json()
    logged_in.put(f"/api/admin/uploads/{state['id']}/chunks/0", content=content[:1000])
    # The browser reconnects and asks what is already there.
    resumed = logged_in.get(f"/api/admin/uploads/{state['id']}").json()
    assert resumed["received_chunks"] == [0]
    too_short = logged_in.put(f"/api/admin/uploads/{state['id']}/chunks/1", content=b"x")
    assert too_short.status_code == 400
    incomplete = logged_in.post(f"/api/admin/uploads/{state['id']}/complete")
    assert incomplete.status_code == 409


def test_upload_refuses_other_file_types(logged_in: TestClient) -> None:
    wrong_name = logged_in.post(
        "/api/admin/uploads", json={"filename": "photo.png", "size_bytes": 10}
    )
    assert wrong_name.status_code == 415
    uploads.CHUNK_SIZE = 1000
    content = b"\x89PNG not really an export" * 10
    state = logged_in.post(
        "/api/admin/uploads", json={"filename": "fake.xlsx", "size_bytes": len(content)}
    ).json()
    for index in range(state["chunk_count"]):
        logged_in.put(
            f"/api/admin/uploads/{state['id']}/chunks/{index}",
            content=content[index * 1000 : (index + 1) * 1000],
        )
    assert logged_in.post(f"/api/admin/uploads/{state['id']}/complete").status_code == 415


def test_full_run_gives_checks_dashboard_and_reports(logged_in: TestClient) -> None:
    data_file = upload(logged_in, demo_workbook(), "stock.xlsx")
    run = start_run(logged_in, data_file["id"])

    assert run["status"] == "succeeded", (run["message"], run["errors"])
    assert run["progress"] == 100
    assert run["checks_total"] == 3 and run["checks_passed"] == 3
    assert run["inputs"]["stock"]["sha256"] == data_file["sha256"]
    assert run["analysis_version"] == "1.0.0"
    kpis = {k["label"]: k["value"] for k in run["dashboard"]["kpis"]}
    assert kpis["Stock value"] == pytest.approx(686000.5)
    assert kpis["Lines with negative stock"] == 1
    by_plant = next(t for t in run["dashboard"]["tables"] if t["id"] == "by-plant")
    assert [r["Plant"] for r in by_plant["rows"]] == ["P100", "P200", "P300"]

    # The demo has no layout of its own: the engine adds the standard dashboard.
    assert {r["key"] for r in run["reports"]} == {"workbook", "report", "slides", "dashboard"}
    dashboard = logged_in.get(f"/api/admin/runs/{run['id']}/reports/dashboard")
    assert dashboard.status_code == 200
    assert dashboard.headers["content-type"].startswith("text/html")
    workbook = logged_in.get(f"/api/admin/runs/{run['id']}/reports/workbook")
    assert workbook.status_code == 200
    sheets = load_workbook(io.BytesIO(workbook.content)).sheetnames
    assert sheets[0] == "Summary" and "Checks" in sheets
    report = logged_in.get(f"/api/admin/runs/{run['id']}/reports/report")
    assert "Stock value by plant" in docx_texts(report.content)

    table = logged_in.get(f"/api/admin/runs/{run['id']}/tables/lines?limit=2").json()
    assert table["total"] == 5 and len(table["rows"]) == 2


def test_same_inputs_give_the_same_result(logged_in: TestClient) -> None:
    first_file = upload(logged_in, demo_workbook(), "a.xlsx")
    second_file = upload(logged_in, demo_csv(), "a.csv")  # same data, other format
    hashes = {
        start_run(logged_in, first_file["id"])["result_hash"],
        start_run(logged_in, first_file["id"])["result_hash"],
        start_run(logged_in, second_file["id"], sheet="CSV")["result_hash"],
    }
    assert len(hashes) == 1


def test_a_wrong_file_fails_with_clear_messages(logged_in: TestClient) -> None:
    rows = [list(r) for r in DEMO_ROWS]
    rows[2][4] = "ninety thousand"
    data_file = upload(logged_in, demo_workbook(rows=rows), "stock.xlsx")

    run = start_run(logged_in, data_file["id"])

    assert run["status"] == "failed"
    assert run["message"] == "Some input files are not right. Nothing was calculated."
    assert run["errors"] == [
        "stock.xlsx, sheet 'Stock', row 6, column 'Value': 'ninety thousand' is not a valid number"
    ]
    assert run["reports"] == []


def test_choosing_a_sheet_without_the_columns_fails(logged_in: TestClient) -> None:
    data_file = upload(logged_in, demo_workbook(), "stock.xlsx")
    run = start_run(logged_in, data_file["id"], sheet="Notes")
    assert run["status"] == "failed"
    assert any("column 'Plant': this column is missing" in e for e in run["errors"])


def test_run_creation_validates_choices(logged_in: TestClient) -> None:
    response = logged_in.post(
        "/api/admin/runs",
        json={"analysis": "demo_stock_summary", "parameters": {}, "inputs": {}},
    )
    assert response.status_code == 400
    assert "Choose a file" in response.json()["detail"]


def test_compare_two_runs(logged_in: TestClient) -> None:
    march = upload(logged_in, demo_workbook(rows=DEMO_ROWS[:2]), "march.xlsx")
    june = upload(logged_in, demo_workbook(), "june.xlsx")
    a = start_run(logged_in, march["id"], label="End of March")
    b = start_run(logged_in, june["id"], label="End of June")

    comparison = logged_in.get(f"/api/admin/runs/compare?a={a['id']}&b={b['id']}").json()

    value = next(m for m in comparison["metrics"] if m["label"] == "Stock value")
    assert value["a"] == pytest.approx(610000.5)
    assert value["b"] == pytest.approx(686000.5)
    assert value["difference"] == pytest.approx(76000)
    assert value["change"] == pytest.approx(76000 / 610000.5)


def test_run_again_with_the_same_files_and_parameters(
    logged_in: TestClient, monkeypatch: pytest.MonkeyPatch
) -> None:
    data_file = upload(logged_in, demo_workbook(), "stock.xlsx")
    first = start_run(logged_in, data_file["id"])
    # The rules of the analysis change: the run again uses the new version.
    from app.analyses.registry import get_analysis

    analysis = get_analysis("demo_stock_summary")
    assert analysis is not None
    monkeypatch.setattr(type(analysis), "version", "9.9.9")
    response = logged_in.post(f"/api/admin/runs/{first['id']}/rerun")
    assert response.status_code == 201, response.text
    again = logged_in.get(f"/api/admin/runs/{response.json()['id']}").json()
    assert again["id"] != first["id"]
    assert again["status"] == "succeeded"
    assert again["analysis_version"] == "9.9.9"
    assert (again["label"], again["parameters"]) == (first["label"], first["parameters"])
    assert again["inputs"] == first["inputs"]
    assert again["result_hash"] == first["result_hash"]
    assert logged_in.post("/api/admin/runs/999999/rerun").status_code == 404


def test_an_input_can_take_several_files(
    logged_in: TestClient, monkeypatch: pytest.MonkeyPatch
) -> None:
    import dataclasses

    from app.analyses.registry import get_analysis

    analysis = get_analysis("demo_stock_summary")
    assert analysis is not None
    spec = dataclasses.replace(analysis.inputs[0], multiple=True)
    monkeypatch.setattr(type(analysis), "inputs", (spec,))
    june = upload(logged_in, demo_csv(), "june.csv")
    july_rows = "Plant;Material;Material Description;Quantity;Value\nP400;M-0009;Masks;5;500\n"
    july = upload(logged_in, july_rows.encode(), "july.csv")
    copy = upload(logged_in, demo_csv(), "june-copy.csv")  # the same bytes as june
    body = {
        "analysis": "demo_stock_summary",
        "label": "Two months",
        "parameters": {"period_end": "2026-07-31"},
        "inputs": {
            "stock": [
                {"data_file_id": june["id"], "sheet": "CSV"},
                {"data_file_id": july["id"], "sheet": "CSV"},
                {"data_file_id": copy["id"], "sheet": "CSV"},
            ]
        },
    }
    created = logged_in.post("/api/admin/runs", json=body)
    assert created.status_code == 201, created.text
    run = logged_in.get(f"/api/admin/runs/{created.json()['id']}").json()
    assert run["status"] == "succeeded", run["errors"]
    stock = run["inputs"]["stock"]
    assert stock["filename"] == "june.csv"
    assert [m["filename"] for m in stock["more"]] == ["july.csv"]
    assert stock["skipped"] == ["june-copy.csv"]
    figures = {m["label"]: m["value"] for m in run["metrics"]}
    assert figures["Lines"] == len(DEMO_ROWS) + 1  # the copy is not counted twice
    assert figures["Stock value"] == pytest.approx(686000.5 + 500)
    # Run again keeps both files.
    again = logged_in.post(f"/api/admin/runs/{run['id']}/rerun").json()
    assert [m["filename"] for m in again["inputs"]["stock"]["more"]] == ["july.csv"]
    # An input that takes one file refuses a list of two.
    monkeypatch.setattr(type(analysis), "inputs", (dataclasses.replace(spec, multiple=False),))
    refused = logged_in.post("/api/admin/runs", json=body)
    assert refused.status_code == 400
    assert "takes one file" in refused.text


def test_cancel_and_delete(logged_in: TestClient, client: TestClient) -> None:
    data_file = upload(logged_in, demo_workbook(), "stock.xlsx")
    run = start_run(logged_in, data_file["id"])
    assert logged_in.post(f"/api/admin/runs/{run['id']}/cancel").json()["status"] == "succeeded"
    assert logged_in.delete(f"/api/admin/runs/{run['id']}").status_code == 204
    assert logged_in.get(f"/api/admin/runs/{run['id']}").status_code == 404
    assert logged_in.delete(f"/api/admin/data-files/{data_file['id']}").status_code == 204


def test_reports_use_the_active_word_template(logged_in: TestClient) -> None:
    template = Document()
    template.add_paragraph("ORGANISATION HEADER")
    template.add_paragraph("{{p body }}")
    buffer = io.BytesIO()
    template.save(buffer)
    uploaded = logged_in.post(
        "/api/admin/templates", files={"file": ("RMS report.docx", buffer.getvalue())}
    ).json()
    assert uploaded["kind"] == "docx"
    logged_in.post(f"/api/admin/templates/{uploaded['id']}/activate")

    run = start_run(logged_in, upload(logged_in, demo_workbook(), "s.xlsx")["id"])

    report = logged_in.get(f"/api/admin/runs/{run['id']}/reports/report")
    assert docx_texts(report.content)[0] == "ORGANISATION HEADER"


def test_a_template_with_banned_wording_makes_the_run_fail(logged_in: TestClient) -> None:
    template = Document()
    template.add_paragraph("Draft prepared with AI help")
    template.add_paragraph("{{p body }}")
    buffer = io.BytesIO()
    template.save(buffer)
    uploaded = logged_in.post(
        "/api/admin/templates", files={"file": ("bad.docx", buffer.getvalue())}
    ).json()
    logged_in.post(f"/api/admin/templates/{uploaded['id']}/activate")

    run = start_run(logged_in, upload(logged_in, demo_workbook(), "s.xlsx")["id"])

    assert run["status"] == "failed"
    assert "banned wording" in run["message"]


def test_templates_must_be_word_or_powerpoint(logged_in: TestClient) -> None:
    response = logged_in.post("/api/admin/templates", files={"file": ("x.docx", b"not a zip")})
    assert response.status_code == 415


def test_data_endpoints_require_login(client: TestClient) -> None:
    for method, path in [
        ("post", "/api/admin/uploads"),
        ("get", "/api/admin/data-files"),
        ("get", "/api/admin/analyses"),
        ("post", "/api/admin/runs"),
        ("get", "/api/admin/runs"),
        ("post", "/api/admin/runs/1/rerun"),
        ("get", "/api/admin/runs/1/reports/workbook"),
        ("get", "/api/admin/templates"),
    ]:
        assert client.request(method, path).status_code == 401, path


def test_analyses_are_listed_with_their_inputs(logged_in: TestClient) -> None:
    analyses = {a["key"]: a for a in logged_in.get("/api/admin/analyses").json()}
    assert set(analyses) == {
        "demo_stock_summary",
        "closing_stock_wbs",
        "expiry_monitoring",
        "rbc_accountability",
        "count_variance",
        "stock_count_reconciliation",
        "stock_sufficiency_amc",
        "order_fill_rate",
        "expiry_risk",
        "stock_in_transit",
        "po_follow_up",
        "warehouse_activity",
        "picker_accuracy",
        "spot_check_accuracy",
    }
    demo = analyses["demo_stock_summary"]
    assert demo["inputs"][0]["dataset"] == "demo_stock_balance"
    assert "period_end" in demo["parameters_schema"]["properties"]


def test_a_template_without_body_placeholder_fails_with_a_clear_message(
    logged_in: TestClient,
) -> None:
    template = Document()
    template.add_paragraph("Header only")
    buffer = io.BytesIO()
    template.save(buffer)
    uploaded = logged_in.post(
        "/api/admin/templates", files={"file": ("no-body.docx", buffer.getvalue())}
    ).json()
    logged_in.post(f"/api/admin/templates/{uploaded['id']}/activate")

    run = start_run(logged_in, upload(logged_in, demo_workbook(), "s.xlsx")["id"])

    assert run["status"] == "failed"
    assert run["message"].startswith("The active report template cannot be used")


def _weekly_workbook() -> bytes:
    """Fake weekly report: one sheet per programme, a title above the header."""
    book = Workbook()
    header = [
        "Product code",
        "Product description",
        "RMS central warehouse",
        "RMS Branches",
        "SDPs",
        "AMC",
    ]
    sheets = {
        "ARV": [["NPC1", "Drug A", 60, 0, 0, 10], ["NPC2", "Drug B", 5, 0, 0, 10]],
        "MCCH": [["NPC3", "Drug C", 80, 0, 0, 10]],
    }
    book.remove(book.active)  # type: ignore[arg-type]
    for name, rows in sheets.items():
        sheet = book.create_sheet(name)
        sheet.append(["WEEKLY STOCK MONITORING", None])
        sheet.append([])
        sheet.append(header)
        for row in rows:
            sheet.append(row)
    book.create_sheet("Notes").append(["Prepared by the planning team"])
    out = io.BytesIO()
    book.save(out)
    return out.getvalue()


def test_an_input_can_take_every_recognised_sheet(logged_in: TestClient) -> None:
    weekly = upload(logged_in, _weekly_workbook(), "weekly.xlsx")
    stock_csv = (
        b"Material,NPC CODE,Material Description,Plant,Unrestricted,Quality Inspection,Blocked,"
        b"WBS Element,Unit Cost\n"
        b"1001,NPC1,Drug A,1000,10,0,0,RMS-PRG-FND-CP,2\n"
        b"1003,NPC3,Drug C,1000,5,0,0,RMS-PRG-FND-CP,4\n"
    )
    stock = upload(logged_in, stock_csv, "ZSTOCKREPORT.csv")
    analyses = {a["key"]: a for a in logged_in.get("/api/admin/analyses").json()}
    weekly_input = next(
        i for i in analyses["stock_sufficiency_amc"]["inputs"] if i["key"] == "weekly"
    )
    assert weekly_input["all_sheets"] is True
    run = start_run(
        logged_in,
        weekly["id"],
        analysis="stock_sufficiency_amc",
        parameters={"report_date": "2026-07-29"},
        inputs={
            "weekly": {"data_file_id": weekly["id"], "sheet": "*"},
            "stock": {"data_file_id": stock["id"], "sheet": stock["sheets"][0]["name"]},
        },
    )
    assert run["status"] == "succeeded", (run["message"], run["errors"])
    kpis = {k["label"]: k["value"] for k in run["dashboard"]["kpis"]}
    assert kpis["Products"] == 3  # two ARV lines and one MCCH line; the Notes sheet is skipped
    summary = next(t for t in run["dashboard"]["tables"] if t["id"] == "summary")
    assert [r["Program"] for r in summary["rows"]] == ["ARV and Opportunistic Infection", "MCCH"]
    assert run["checks_passed"] == run["checks_total"]


def test_file_list_counts_the_runs_that_used_each_file(logged_in: TestClient) -> None:
    used = upload(logged_in, demo_workbook(), "used.xlsx")
    unused = upload(logged_in, demo_workbook(rows=DEMO_ROWS[:2]), "unused.xlsx")
    start_run(logged_in, used["id"])
    start_run(logged_in, used["id"])
    files = {f["id"]: f for f in logged_in.get("/api/admin/data-files").json()}
    assert files[used["id"]]["used_by_runs"] == 2
    assert files[unused["id"]]["used_by_runs"] == 0
    assert logged_in.get(f"/api/admin/data-files/{used['id']}").json()["used_by_runs"] == 2
