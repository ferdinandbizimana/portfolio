"""Every run gives an Excel workbook and an HTML dashboard."""

import io

import pandas as pd
from openpyxl import load_workbook

from app.analyses.base import Check
from app.analyses.dashboard import Kpi
from app.reporting.documents import ReportFile
from app.reporting.guard import check_report
from app.reporting.interactive import read_payload
from app.reporting.standard import ensure_downloads

TABLES = {
    "by_plant": pd.DataFrame({"Plant": ["P1", "P2"], "Value": [10.5, 20.0], "Lines": [1, 2]}),
    "empty": pd.DataFrame({"A": []}),
}
METRICS = [Kpi(label="Total value", value=30.5, format="money"), Kpi(label="Lines", value=3)]
CHECKS = [Check("total", "Lines add up", expected=3, actual=3)]


def downloads(reports: list[ReportFile]) -> list[ReportFile]:
    return ensure_downloads(
        reports,
        title="Fake analysis",
        subtitle="Invented data",
        name="Fake_analysis_run_1",
        tables=TABLES,
        metrics=METRICS,
        checks=CHECKS,
    )


def test_missing_workbook_and_dashboard_are_added() -> None:
    files = downloads([])
    assert [(f.key, f.kind) for f in files] == [("workbook", "xlsx"), ("dashboard", "html")]
    for report in files:
        check_report(report.kind, report.content)
    workbook = load_workbook(io.BytesIO(files[0].content))
    assert len(workbook.sheetnames) >= 2  # summary and the non-empty table
    data = read_payload(files[1].content)
    assert data is not None
    assert set(data["facts"]) == {"key_figures", "t1"}  # the empty table is left out
    assert data["facts"]["t1"]["rows"] == 2


def test_an_analysis_with_its_own_files_keeps_them() -> None:
    own = [
        ReportFile("workbook", "Excel workbook", "a.xlsx", "xlsx", b"x"),
        ReportFile("dashboard", "Interactive dashboard (HTML)", "a.html", "html", b"y"),
    ]
    assert downloads(own) == own
