import datetime as dt
from pathlib import Path

import pandas as pd
import pyarrow.parquet as pq
import pytest

from app import datasets
from app.ingest.parse import ParseError, cell_text, detect_format, parse_file
from tests.samples import DEMO_HEADERS, demo_csv, demo_workbook


def read(path: Path) -> pd.DataFrame:
    frame: pd.DataFrame = pq.read_table(path).to_pandas()
    return frame


def test_cell_text_keeps_values_exact() -> None:
    assert cell_text(1200.0) == "1200"
    assert cell_text(250000.5) == "250000.5"
    assert cell_text(0.1 + 0.2) == "0.30000000000000004"
    assert cell_text(dt.datetime(2026, 6, 30)) == "2026-06-30"
    assert cell_text(dt.datetime(2026, 6, 30, 14, 5)) == "2026-06-30 14:05:00"
    assert cell_text("  M-1 ") == "M-1"
    assert cell_text("") is None
    assert cell_text(None) is None


def test_excel_sheets_are_found_below_title_lines(tmp_path: Path) -> None:
    source = tmp_path / "export.xlsx"
    source.write_bytes(demo_workbook(title_rows=3))

    sheets = parse_file(source, tmp_path / "out")

    stock, notes = sheets
    assert stock.name == "Stock"
    assert stock.header_row == 4
    assert stock.columns == DEMO_HEADERS
    assert stock.row_count == 5  # trailing blank line ignored
    frame = read(tmp_path / "out" / stock.parquet_path)
    assert frame.loc[0, "Material"] == "M-0001"
    assert frame.loc[1, "Value"] == "250000.5"
    assert frame.loc[4, "Material Description"] is None
    assert notes.name == "Notes"


def test_csv_with_semicolons_and_windows_encoding(tmp_path: Path) -> None:
    source = tmp_path / "export.csv"
    source.write_bytes(demo_csv(separator=";", encoding="cp1252").replace(b"Gloves", b"Gants \xe9"))

    [sheet] = parse_file(source, tmp_path / "out")

    frame = read(tmp_path / "out" / sheet.parquet_path)
    assert sheet.columns == DEMO_HEADERS
    assert str(frame.loc[0, "Material Description"]).startswith("Gants é")


def test_duplicate_and_empty_headers_get_unique_names(tmp_path: Path) -> None:
    source = tmp_path / "dup.xlsx"
    source.write_bytes(
        demo_workbook(headers=["Plant", "Plant", "Material", None, "Value"], title_rows=0)  # type: ignore[list-item]
    )
    sheet = parse_file(source, tmp_path / "out")[0]
    assert sheet.columns == ["Plant", "Plant (2)", "Material", "Column 4", "Value"]


def test_unknown_file_type_is_refused(tmp_path: Path) -> None:
    source = tmp_path / "x.xlsx"
    source.write_bytes(b"\x00\x01\x02 not a spreadsheet")
    assert detect_format(source.read_bytes()) is None
    with pytest.raises(ParseError, match="not an Excel"):
        parse_file(source, tmp_path / "out")


def test_matching_recognises_the_demo_dataset() -> None:
    matches = datasets.match_sheet(DEMO_HEADERS, "Stock")
    assert [m.dataset for m in matches] == ["demo_stock_balance"]
    # Header spelling and case do not matter; the optional column may be absent.
    assert datasets.match_sheet(["PLANT", "material", "quantity", "Value "], "Sheet1")
    assert datasets.match_sheet(["Plant", "Material", "Quantity"], "Sheet1") == []


def raw_frame(rows: list[list[str | None]]) -> pd.DataFrame:
    return pd.DataFrame(rows, columns=DEMO_HEADERS)


def test_coerce_converts_types() -> None:
    spec = datasets.get("demo_stock_balance")
    coerced = datasets.coerce(
        spec,
        raw_frame([["P1", "M1", None, "1,200", "15000-"], ["P2", "M2", "x", "-3.5", "0"]]),
        header_row=1,
    )
    assert coerced.ok
    assert coerced.frame["Quantity"].tolist() == [1200.0, -3.5]
    assert coerced.frame["Value"].tolist() == [-15000.0, 0.0]
    assert coerced.frame["Material Description"].tolist() == [None, "x"]


def test_coerce_names_the_row_and_column_of_bad_values() -> None:
    spec = datasets.get("demo_stock_balance")
    coerced = datasets.coerce(spec, raw_frame([["P1", "M1", None, "12", "abc"]]), header_row=3)
    [issue] = coerced.issues
    assert issue.text("stock.xlsx", "Stock") == (
        "stock.xlsx, sheet 'Stock', row 4, column 'Value': 'abc' is not a valid number"
    )


def test_coerce_reports_missing_columns() -> None:
    spec = datasets.get("demo_stock_balance")
    frame = pd.DataFrame([["P1", "M1", "3"]], columns=["Plant", "Material", "Quantity"])
    [issue] = datasets.coerce(spec, frame, header_row=1).issues
    assert (
        issue.text("f.csv", "CSV") == "f.csv, sheet 'CSV', column 'Value': this column is missing"
    )


def test_sheets_without_a_table_are_left_out(tmp_path: Path) -> None:
    from openpyxl import Workbook

    workbook = Workbook()
    notes = workbook.active
    assert notes is not None
    notes.title = "Notes"
    notes.append(["Prepared for the quarterly review"])
    notes.append(["Stock value", 1000])  # label and value pairs: no table
    data = workbook.create_sheet("Stock")
    data.append(DEMO_HEADERS)
    data.append(["P100", "M-1", "Gloves", 1, 2])
    path = tmp_path / "mixed.xlsx"
    workbook.save(path)

    assert [s.name for s in parse_file(path, tmp_path / "out")] == ["Stock"]

    only_notes = Workbook()
    sheet = only_notes.active
    assert sheet is not None
    sheet.append(["Stock value", 1000])
    path = tmp_path / "notes.xlsx"
    only_notes.save(path)
    with pytest.raises(ParseError, match="sheet 'Sheet': No header row"):
        parse_file(path, tmp_path / "out2")
