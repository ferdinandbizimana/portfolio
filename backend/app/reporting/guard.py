"""Language guard: generated reports must never contain banned wording.

Ferdinand's rule: no "AI" and no "to be confirmed" in any report. "AI" is
matched as a whole word in capitals (so "AID", "maintain" or "Kigali" pass);
"to be confirmed" in any case.
"""

import html
import io
import re
from collections.abc import Iterable

from docx import Document as load_docx
from openpyxl import load_workbook
from pptx import Presentation

from app.reporting.interactive import unpack

BANNED = (
    (re.compile(r"\bAI\b"), "AI"),
    (re.compile(r"to be confirmed", re.IGNORECASE), "to be confirmed"),
)


class BannedWordingError(Exception):
    pass


def find_banned(texts: Iterable[str]) -> list[str]:
    """Each banned phrase found, with the text around it."""
    found: list[str] = []
    for text in texts:
        for pattern, name in BANNED:
            for match in pattern.finditer(text):
                start = max(match.start() - 30, 0)
                found.append(f'"{name}" in: …{text[start : match.end() + 30]}…')
    return found


def docx_texts(content: bytes) -> list[str]:
    document = load_docx(io.BytesIO(content))
    texts = [p.text for p in document.paragraphs]
    for table in document.tables:
        texts += [cell.text for row in table.rows for cell in row.cells]
    for section in document.sections:
        texts += [p.text for p in section.header.paragraphs]
        texts += [p.text for p in section.footer.paragraphs]
    return texts


def pptx_texts(content: bytes) -> list[str]:
    texts: list[str] = []
    for slide in Presentation(io.BytesIO(content)).slides:
        for shape in slide.shapes:
            if shape.has_text_frame:
                texts.append(shape.text_frame.text)
            if getattr(shape, "has_table", False) and shape.has_table:
                texts += [cell.text for row in shape.table.rows for cell in row.cells]
            if getattr(shape, "has_chart", False) and shape.has_chart:
                chart = shape.chart
                if chart.has_title:
                    texts.append(chart.chart_title.text_frame.text)
        if slide.has_notes_slide:
            texts.append(slide.notes_slide.notes_text_frame.text)
    return texts


def xlsx_texts(content: bytes) -> list[str]:
    workbook = load_workbook(io.BytesIO(content), read_only=True)
    texts = list(workbook.sheetnames)
    for sheet in workbook.worksheets:
        for row in sheet.iter_rows(values_only=True):
            texts += [v for v in row if isinstance(v, str) and not v.startswith("=")]
    workbook.close()
    return texts


def _strings(value: object) -> list[str]:
    if isinstance(value, str):
        return [value]
    if isinstance(value, dict):
        return [s for k, v in value.items() for s in [*_strings(k), *_strings(v)]]
    if isinstance(value, list):
        return [s for v in value for s in _strings(v)]
    return []


def html_texts(content: bytes) -> list[str]:
    """Visible text of an interactive dashboard and every text in its data."""
    page = content.decode("utf-8")
    data = re.search(r'<script id="data"[^>]*>(.*?)</script>', page, re.S)
    visible = re.sub(r"<script[^>]*>.*?</script>|<style[^>]*>.*?</style>", " ", page, flags=re.S)
    texts = [html.unescape(t) for t in re.split(r"<[^>]+>", visible) if t.strip()]
    if data:
        texts += _strings(unpack(data.group(1)))
    return texts


def check_report(kind: str, content: bytes) -> None:
    extract = {"docx": docx_texts, "pptx": pptx_texts, "xlsx": xlsx_texts, "html": html_texts}[kind]
    found = find_banned(extract(content))
    if found:
        raise BannedWordingError("; ".join(found[:5]))
