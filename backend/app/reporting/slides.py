"""PowerPoint decks: the RMS management (SMT) layout, or an uploaded template.

Built-in layout (no template), after Ferdinand's "Order Fill Rate FY2025/26"
SMT deck and the RMS Brand Guide: a blue cover with the white logo; content
slides with a blue title bar, the logo top right on a white tab, the key
message in italics under the bar, native (editable) charts next to their table,
a "So what" box, and a footer line with the report, the period and the page;
then recommendations (owner and timing), data notes, the checks and a closing
slide with the tagline.

With an uploaded template, its slide layouts are used ("Title Slide" for the
cover and "Title and Content" / "Title Only" for the rest, found by name, else
by position) and the template's own example slides are removed.
"""

import io
import math
from collections.abc import Sequence
from pathlib import Path
from typing import Any

import pandas as pd
from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE, XL_LABEL_POSITION, XL_LEGEND_POSITION
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Emu, Inches, Pt

from app.reporting import brand
from app.reporting.documents import (
    ChartBlock,
    Highlight,
    KpiCard,
    Recommendation,
    ReportDocument,
    Section,
    TableBlock,
    ValueFormat,
    format_value,
)

MAX_TABLE_ROWS = 12

_CHART_TYPES = {
    "bar": XL_CHART_TYPE.BAR_CLUSTERED,
    "column": XL_CHART_TYPE.COLUMN_CLUSTERED,
    "pie": XL_CHART_TYPE.PIE,
    "line": XL_CHART_TYPE.LINE_MARKERS,
    "stacked": XL_CHART_TYPE.COLUMN_STACKED,
}

SLIDE_W, SLIDE_H = Inches(13.333), Inches(7.5)
MARGIN = Inches(0.4)
BAR_H = Inches(0.72)
CONTENT_TOP = Inches(1.35)
FOOTER_TOP = Inches(7.08)
CONTENT_BOTTOM = Inches(6.95)


def _rgb(colour: str) -> RGBColor:
    return RGBColor(*brand.rgb(colour))  # type: ignore[no-untyped-call]


# --- Uploaded template (layouts by name) -----------------------------------------------


def _layout(prs: Any, name: str, fallback: int) -> Any:
    for layout in prs.slide_layouts:
        if layout.name.lower() == name.lower():
            return layout
    layouts = list(prs.slide_layouts)
    return layouts[min(fallback, len(layouts) - 1)]


def _remove_existing_slides(prs: Any) -> None:
    slide_ids = prs.slides._sldIdLst
    for slide_id in list(slide_ids):
        prs.part.drop_rel(slide_id.rId)
        slide_ids.remove(slide_id)


def _content_box(prs: Any, slide: Any) -> tuple[int, int, int, int]:
    """Area below the title for tables and charts (in EMU)."""
    top = int(Inches(1.5))
    if slide.shapes.title is not None:
        title = slide.shapes.title
        top = int(title.top + title.height + Inches(0.2))
    left = int(Inches(0.6))
    width = int(prs.slide_width - 2 * left)
    height = int(prs.slide_height - top - Inches(0.5))
    return left, top, width, height


def _title_only(prs: Any, title: str) -> Any:
    slide = prs.slides.add_slide(_layout(prs, "Title Only", 5))
    if slide.shapes.title is not None:
        slide.shapes.title.text = title
    return slide


def _text_slide(prs: Any, title: str, lines: Sequence[str]) -> None:
    slide = prs.slides.add_slide(_layout(prs, "Title and Content", 1))
    if slide.shapes.title is not None:
        slide.shapes.title.text = title
    body = next((p for p in slide.placeholders if p.placeholder_format.idx == 1), None)
    if body is not None and lines:
        frame = body.text_frame
        frame.text = lines[0]
        for line in lines[1:]:
            frame.add_paragraph().text = line


def _template_deck(prs: Any, document: ReportDocument) -> None:
    cover = prs.slides.add_slide(_layout(prs, "Title Slide", 0))
    if cover.shapes.title is not None:
        cover.shapes.title.text = document.title
    subtitle = next((p for p in cover.placeholders if p.placeholder_format.idx == 1), None)
    if subtitle is not None:
        subtitle.text = f"{document.subtitle}\n{document.date_text}"
    if document.summary_message or document.kpis or document.highlights:
        _text_slide(
            prs,
            "Executive summary",
            [
                *([document.summary_message] if document.summary_message else []),
                *(
                    f"{k.label}: {k.value}" + (f" ({k.note})" if k.note else "")
                    for k in document.kpis
                ),
                *(f"{h.title}: {h.text}" for h in document.highlights),
            ],
        )
    for section in document.sections:
        lines = [
            *([section.message] if section.message else []),
            *section.paragraphs,
            *section.bullets,
            *([f"So what: {section.takeaway}"] if section.takeaway else []),
        ]
        if lines:
            _text_slide(prs, section.heading, lines)
        for chart in section.charts:
            slide = _title_only(prs, chart.title)
            _add_chart(slide, chart, *_content_box(prs, slide))
        for block in section.tables:
            slide = _title_only(prs, block.title)
            left, top, width, height = _content_box(prs, slide)
            _add_table(slide, block, left, top, width, height)
    if document.recommendations:
        _text_slide(
            prs,
            "Recommendations and next steps",
            [f"{r.action} ({r.owner}; {r.timing})" for r in document.recommendations],
        )
    if document.notes:
        _text_slide(prs, "Data notes", [f"{n.title}: {n.text}" for n in document.notes])
    if document.checks:
        passed = sum(c.passed for c in document.checks)
        _text_slide(
            prs,
            "Report checks",
            [
                f"{passed} of {len(document.checks)} checks passed.",
                *(f"{'Passed' if c.passed else 'Failed'}: {c.label}" for c in document.checks),
            ],
        )


# --- Drawing helpers --------------------------------------------------------------------


def _box(
    slide: Any,
    left: int,
    top: int,
    width: int,
    height: int,
    fill: str | None,
    line: str | None = None,
    shape: Any = MSO_SHAPE.RECTANGLE,
) -> Any:
    item = slide.shapes.add_shape(shape, left, top, width, height)
    if fill is None:
        item.fill.background()
    else:
        item.fill.solid()
        item.fill.fore_color.rgb = _rgb(fill)
    if line is None:
        item.line.fill.background()
    else:
        item.line.color.rgb = _rgb(line)
        item.line.width = Pt(0.75)
    item.shadow.inherit = False
    if shape == MSO_SHAPE.ROUNDED_RECTANGLE:
        item.adjustments[0] = 0.08
    return item


def _text(
    slide: Any,
    left: int,
    top: int,
    width: int,
    height: int,
    text: str,
    size: float = 11,
    colour: str = brand.INK,
    bold: bool = False,
    italic: bool = False,
    font: str = brand.BODY_FONT,
    align: Any = PP_ALIGN.LEFT,
    anchor: Any = MSO_ANCHOR.TOP,
) -> Any:
    item = slide.shapes.add_textbox(left, top, width, height)
    frame = item.text_frame
    frame.word_wrap = True
    frame.vertical_anchor = anchor
    frame.margin_left = frame.margin_right = Inches(0.08)
    frame.margin_top = frame.margin_bottom = Inches(0.04)
    for k, line in enumerate(text.split("\n")):
        paragraph = frame.paragraphs[0] if k == 0 else frame.add_paragraph()
        paragraph.alignment = align
        run = paragraph.add_run()
        run.text = line
        run.font.size = Pt(size)
        run.font.bold = bold
        run.font.italic = italic
        run.font.name = font
        run.font.color.rgb = _rgb(colour)
    return item


def _logo(slide: Any, variant: str, left: int, top: int, height: int) -> int:
    """Place the logo at its proportions; returns its width (0 when there is no logo)."""
    content, size = brand.logo(variant), brand.logo_size(variant)
    if content is None or size is None:
        return 0
    width = int(height * size[0] / size[1])
    slide.shapes.add_picture(io.BytesIO(content), left, top, width, height)
    return width


def _accent(slide: Any) -> None:
    """The guide's corner stripes: blue and a narrow yellow bar, bottom right."""
    _box(slide, SLIDE_W - Inches(1.6), SLIDE_H - Inches(0.1), Inches(1.0), Inches(0.1), brand.BLUE)
    _box(
        slide, SLIDE_W - Inches(0.6), SLIDE_H - Inches(0.1), Inches(0.6), Inches(0.1), brand.YELLOW
    )


class Deck:
    def __init__(self, document: ReportDocument) -> None:
        self.document = document
        self.prs = Presentation()
        self.prs.slide_width, self.prs.slide_height = SLIDE_W, SLIDE_H
        self.blank = self.prs.slide_layouts[6]
        self.footer = brand.footer(document.title, document.date_text, document.draft)

    def slide(self, title: str, message: str = "") -> Any:
        slide = self.prs.slides.add_slide(self.blank)
        _box(slide, 0, 0, SLIDE_W, BAR_H, brand.BLUE)
        logo_h = Inches(0.5)
        size = brand.logo_size("colour")
        if size is not None:
            tab_w = int(logo_h * size[0] / size[1]) + Inches(0.24)
            left = SLIDE_W - MARGIN - tab_w
            _box(
                slide,
                left,
                Inches(0.07),
                tab_w,
                BAR_H - Inches(0.14),
                brand.WHITE,
                shape=MSO_SHAPE.ROUNDED_RECTANGLE,
            )
            _logo(slide, "colour", left + Inches(0.12), Inches(0.11), logo_h)
        _text(
            slide,
            MARGIN,
            0,
            SLIDE_W - Inches(4.2),
            BAR_H,
            title,
            24,
            brand.WHITE,
            bold=True,
            font=brand.HEADING_FONT,
            anchor=MSO_ANCHOR.MIDDLE,
        )
        if message:
            _text(
                slide,
                MARGIN,
                Inches(0.8),
                SLIDE_W - 2 * MARGIN,
                Inches(0.5),
                message,
                13,
                brand.MUTED,
                italic=True,
            )
        number = len(self.prs.slides)
        _box(slide, MARGIN, FOOTER_TOP - Inches(0.04), SLIDE_W - 2 * MARGIN, Emu(9525), brand.LINE)
        _text(slide, MARGIN, FOOTER_TOP, Inches(10), Inches(0.3), self.footer, 9, brand.MUTED)
        _text(
            slide,
            SLIDE_W - MARGIN - Inches(2.2),
            FOOTER_TOP,
            Inches(1.6),
            Inches(0.3),
            str(number),
            9,
            brand.MUTED,
            align=PP_ALIGN.RIGHT,
        )
        _accent(slide)
        return slide

    # --- Cover and closing -----------------------------------------------------------

    def cover(self) -> None:
        doc = self.document
        slide = self.prs.slides.add_slide(self.blank)
        _box(slide, 0, 0, SLIDE_W, SLIDE_H, brand.BLUE)
        _box(slide, Inches(8.6), 0, SLIDE_W - Inches(8.6), SLIDE_H, brand.tint(brand.BLUE, 0.12))
        _box(slide, Inches(8.6), SLIDE_H - Inches(0.5), Inches(2.6), Inches(0.5), brand.YELLOW)
        _logo(slide, "white", Inches(0.7), Inches(0.6), Inches(0.95))
        _text(
            slide,
            Inches(0.7),
            Inches(2.4),
            Inches(7.6),
            Inches(1.6),
            doc.title.upper(),
            36,
            brand.WHITE,
            bold=True,
            font=brand.HEADING_FONT,
            anchor=MSO_ANCHOR.BOTTOM,
        )
        _box(slide, Inches(0.78), Inches(4.1), Inches(0.9), Inches(0.06), brand.YELLOW)
        _text(
            slide,
            Inches(0.7),
            Inches(4.3),
            Inches(7.6),
            Inches(0.9),
            doc.subtitle,
            16,
            brand.WHITE,
            italic=True,
        )
        _text(
            slide,
            Inches(0.7),
            Inches(5.2),
            Inches(7.6),
            Inches(0.5),
            doc.date_text,
            13,
            brand.WHITE,
        )
        footer = " · ".join(p for p in (doc.audience, "DRAFT" if doc.draft else "") if p)
        if footer:
            _text(
                slide, Inches(0.7), Inches(6.7), Inches(7.6), Inches(0.4), footer, 11, brand.WHITE
            )
        _text(
            slide,
            Inches(8.9),
            Inches(3.2),
            Inches(4.0),
            Inches(1.2),
            brand.TAGLINE,
            22,
            brand.BLUE,
            bold=True,
            font=brand.HEADING_FONT,
        )

    def closing(self) -> None:
        slide = self.prs.slides.add_slide(self.blank)
        _box(slide, 0, 0, SLIDE_W, SLIDE_H, brand.BLUE)
        size = brand.logo_size("white")
        if size is not None:
            height = Inches(1.2)
            width = int(height * size[0] / size[1])
            _logo(slide, "white", (SLIDE_W - width) // 2, Inches(1.9), height)
        _text(
            slide,
            0,
            Inches(3.5),
            SLIDE_W,
            Inches(0.8),
            "Thank You",
            36,
            brand.WHITE,
            bold=True,
            font=brand.HEADING_FONT,
            align=PP_ALIGN.CENTER,
        )
        _text(
            slide,
            0,
            Inches(4.3),
            SLIDE_W,
            Inches(0.5),
            "Questions & Discussion",
            16,
            brand.WHITE,
            italic=True,
            align=PP_ALIGN.CENTER,
        )
        _text(
            slide,
            0,
            Inches(6.6),
            SLIDE_W,
            Inches(0.4),
            brand.TAGLINE,
            14,
            brand.YELLOW,
            bold=True,
            font=brand.HEADING_FONT,
            align=PP_ALIGN.CENTER,
        )

    # --- Building blocks ------------------------------------------------------------

    def kpis(
        self, slide: Any, cards: Sequence[KpiCard], top: int, height: int = Inches(1.3)
    ) -> int:
        cards = list(cards)[:6]
        if not cards:
            return top
        gap = Inches(0.15)
        width = (SLIDE_W - 2 * MARGIN - gap * (len(cards) - 1)) // len(cards)
        for k, card in enumerate(cards):
            accent, _fill = brand.TONES[card.tone]
            left = MARGIN + k * (width + gap)
            _box(slide, left, top, width, height, brand.PANEL)
            _box(slide, left, top + height - Inches(0.07), width, Inches(0.07), accent)
            _text(
                slide,
                left,
                top + Inches(0.08),
                width,
                Inches(0.55),
                card.value,
                24 if len(card.value) <= 7 else 18 if len(card.value) <= 11 else 14,
                accent,
                bold=True,
                font=brand.HEADING_FONT,
                align=PP_ALIGN.CENTER,
                anchor=MSO_ANCHOR.MIDDLE,
            )
            _text(
                slide,
                left,
                top + Inches(0.62),
                width,
                Inches(0.3),
                card.label,
                10,
                brand.INK,
                bold=True,
                align=PP_ALIGN.CENTER,
            )
            if card.note:
                _text(
                    slide,
                    left,
                    top + Inches(0.9),
                    width,
                    Inches(0.3),
                    card.note,
                    8,
                    brand.MUTED,
                    align=PP_ALIGN.CENTER,
                )
        return top + height + Inches(0.2)

    def panels(
        self,
        slide: Any,
        items: Sequence[Highlight],
        left: int,
        top: int,
        width: int,
        height: int,
        columns: int | None = None,
    ) -> None:
        items = list(items)
        if not items:
            return
        columns = columns or min(len(items), 3)
        rows = math.ceil(len(items) / columns)
        gap = Inches(0.15)
        each_w = (width - gap * (columns - 1)) // columns
        each_h = (height - gap * (rows - 1)) // rows
        # Boxes as tall as their text needs (about 13 characters per inch at 11 pt).
        per_line = max(20, int(each_w / Inches(1) * 13))
        lines = max(
            sum(max(1, math.ceil(len(part) / per_line)) for part in item.text.split("\n"))
            for item in items
        )
        each_h = min(each_h, max(Inches(1.0), Inches(0.6) + lines * Inches(0.22)))
        for k, item in enumerate(items):
            accent, fill = brand.TONES[item.tone]
            x = left + (k % columns) * (each_w + gap)
            y = top + (k // columns) * (each_h + gap)
            _box(slide, x, y, each_w, each_h, fill, brand.LINE)
            _box(slide, x, y, Inches(0.06), each_h, accent)
            _text(
                slide,
                x + Inches(0.12),
                y + Inches(0.08),
                each_w - Inches(0.2),
                Inches(0.35),
                item.title,
                12,
                accent,
                bold=True,
                font=brand.HEADING_FONT,
            )
            size = 11 if len(item.text) < 260 else 10 if len(item.text) < 420 else 9
            _text(
                slide,
                x + Inches(0.12),
                y + Inches(0.45),
                each_w - Inches(0.2),
                each_h - Inches(0.5),
                item.text,
                size,
                brand.INK,
            )

    def takeaway(
        self,
        slide: Any,
        text: str,
        left: int,
        top: int,
        width: int,
        height: int,
        title: str = "So what",
    ) -> None:
        self.panels(slide, [Highlight(title, text, "blue")], left, top, width, height, columns=1)

    # --- Slides ---------------------------------------------------------------------

    def executive_summary(self) -> None:
        doc = self.document
        if not (doc.summary_message or doc.kpis or doc.highlights):
            return
        slide = self.slide("Executive summary", doc.summary_message)
        top = self.kpis(slide, doc.kpis, CONTENT_TOP)
        self.panels(slide, doc.highlights, MARGIN, top, SLIDE_W - 2 * MARGIN, CONTENT_BOTTOM - top)

    def section(self, section: Section) -> None:
        visuals: list[ChartBlock | TableBlock] = [*section.charts, *section.tables]
        text = "\n".join([*section.paragraphs, *(f"• {b}" for b in section.bullets)])
        pairs = [visuals[i : i + 2] for i in range(0, len(visuals), 2)] or [[]]
        for index, pair in enumerate(pairs):
            first, last = index == 0, index == len(pairs) - 1
            slide = self.slide(
                section.heading if first else f"{section.heading} (continued)",
                section.message if first else "",
            )
            top = (
                self.kpis(slide, section.kpis, CONTENT_TOP, Inches(1.15)) if first else CONTENT_TOP
            )
            reading = section.takeaway if last else ""
            extra = text if first else ""
            box_h = Inches(1.15) if (reading or extra) and pair else 0
            bottom = CONTENT_BOTTOM - box_h - (Inches(0.15) if box_h else 0)
            width = SLIDE_W - 2 * MARGIN
            if not pair:
                if section.highlights:
                    panel_h = (
                        bottom - top if not (reading or extra) else (CONTENT_BOTTOM - top) * 2 // 3
                    )
                    self.panels(slide, section.highlights, MARGIN, top, width, panel_h)
                    top += panel_h + Inches(0.15)
                body = "\n".join(t for t in (extra, reading) if t)
                if body:
                    self.takeaway(
                        slide,
                        body,
                        MARGIN,
                        top,
                        width,
                        CONTENT_BOTTOM - top,
                        "So what" if reading and not extra else "Reading",
                    )
                continue
            if len(pair) == 1:
                boxes: list[tuple[int, int, int, int]] = [(int(MARGIN), top, width, bottom - top)]
            else:
                left_w = width * 55 // 100
                boxes = [
                    (int(MARGIN), top, left_w - Inches(0.1), bottom - top),
                    (
                        MARGIN + left_w + Inches(0.1),
                        top,
                        width - left_w - Inches(0.1),
                        bottom - top,
                    ),
                ]
            for visual, (x, y, w, h) in zip(pair, boxes, strict=True):
                if isinstance(visual, ChartBlock):
                    _add_chart(slide, visual, x, y, w, h)
                else:
                    _add_table(slide, visual, x, y, w, h, title=True)
            if section.highlights and last:
                more = self.slide(f"{section.heading}: reading", "")
                self.panels(
                    more,
                    section.highlights,
                    MARGIN,
                    CONTENT_TOP,
                    width,
                    CONTENT_BOTTOM - CONTENT_TOP,
                )
            if box_h:
                if extra and reading:
                    half = (width - Inches(0.15)) // 2
                    self.takeaway(
                        slide, extra, MARGIN, bottom + Inches(0.15), half, box_h, "Reading"
                    )
                    self.takeaway(
                        slide,
                        reading,
                        MARGIN + half + Inches(0.15),
                        bottom + Inches(0.15),
                        half,
                        box_h,
                    )
                else:
                    self.takeaway(
                        slide,
                        reading or extra,
                        MARGIN,
                        bottom + Inches(0.15),
                        width,
                        box_h,
                        "So what" if reading else "Reading",
                    )

    def recommendations(self, items: Sequence[Recommendation]) -> None:
        if not items:
            return
        slide = self.slide(
            "Recommendations and next steps",
            "In priority order, with a proposed owner and timing for SMT to assign",
        )
        items = list(items)[:7]
        row_h = min(Inches(0.75), (CONTENT_BOTTOM - CONTENT_TOP) // len(items))
        for k, item in enumerate(items):
            y = CONTENT_TOP + k * row_h
            _box(
                slide,
                MARGIN,
                y + Inches(0.05),
                SLIDE_W - 2 * MARGIN,
                row_h - Inches(0.1),
                brand.PANEL,
            )
            circle = _box(
                slide,
                MARGIN + Inches(0.12),
                y + (row_h - Inches(0.42)) // 2,
                Inches(0.42),
                Inches(0.42),
                brand.BLUE,
                shape=MSO_SHAPE.OVAL,
            )
            frame = circle.text_frame
            frame.text = str(k + 1)
            run = frame.paragraphs[0].runs[0]
            run.font.size, run.font.bold, run.font.color.rgb = Pt(12), True, _rgb(brand.WHITE)
            frame.paragraphs[0].alignment = PP_ALIGN.CENTER
            _text(
                slide,
                MARGIN + Inches(0.7),
                y + Inches(0.05),
                Inches(8.3),
                row_h - Inches(0.1),
                item.action,
                11,
                brand.INK,
                bold=True,
                anchor=MSO_ANCHOR.MIDDLE,
            )
            _text(
                slide,
                Inches(9.6),
                y + Inches(0.05),
                Inches(2.0),
                row_h - Inches(0.1),
                item.owner,
                10,
                brand.BLUE,
                bold=True,
                anchor=MSO_ANCHOR.MIDDLE,
            )
            _text(
                slide,
                Inches(11.6),
                y + Inches(0.05),
                Inches(1.3),
                row_h - Inches(0.1),
                item.timing,
                10,
                brand.INK,
                anchor=MSO_ANCHOR.MIDDLE,
            )

    def notes(self, items: Sequence[Highlight]) -> None:
        if not items:
            return
        slide = self.slide("Data notes", "What the numbers rest on")
        self.panels(
            slide,
            items,
            MARGIN,
            CONTENT_TOP,
            SLIDE_W - 2 * MARGIN,
            CONTENT_BOTTOM - CONTENT_TOP,
            columns=2 if len(items) > 1 else 1,
        )

    def checks(self) -> None:
        checks = self.document.checks
        if not checks:
            return
        passed = sum(c.passed for c in checks)
        slide = self.slide(
            "Report checks",
            f"{passed} of {len(checks)} checks passed: the figures tie back to the source files",
        )
        frame = pd.DataFrame(
            {
                "Check": [c.label for c in checks],
                "Expected": [c.expected for c in checks],
                "Actual": [c.actual for c in checks],
                "Result": ["Passed" if c.passed else "Failed" for c in checks],
            }
        )
        block = TableBlock("Checks", frame, {"Expected": "number", "Actual": "number"})
        _add_table(
            slide, block, MARGIN, CONTENT_TOP, SLIDE_W - 2 * MARGIN, CONTENT_BOTTOM - CONTENT_TOP
        )

    def build(self) -> bytes:
        self.cover()
        self.executive_summary()
        for section in self.document.sections:
            self.section(section)
        self.recommendations(self.document.recommendations)
        self.notes(self.document.notes)
        self.checks()
        self.closing()
        buffer = io.BytesIO()
        self.prs.save(buffer)
        return buffer.getvalue()


# --- Tables and charts (both layouts) ------------------------------------------------


def _numeric(block: TableBlock, column: str) -> bool:
    return block.formats.get(column) in ("money", "number", "integer", "percent")


def _add_table(
    slide: Any, block: TableBlock, left: int, top: int, width: int, height: int, title: bool = False
) -> None:
    if title:
        _text(
            slide,
            left,
            top,
            width,
            Inches(0.32),
            block.title,
            12,
            brand.BLUE,
            bold=True,
            font=brand.HEADING_FONT,
        )
        top += Inches(0.36)
        height -= Inches(0.36)
    limit = min(block.max_rows or MAX_TABLE_ROWS, MAX_TABLE_ROWS)
    frame = block.frame.head(limit)
    columns = [str(c) for c in frame.columns]
    if not columns:
        return
    rows = len(frame) + 1
    row_h = min(Inches(0.32), height // max(rows, 1))
    shape = slide.shapes.add_table(rows, len(columns), left, top, width, row_h * rows)
    table = shape.table
    table.first_row = True
    table.horz_banding = False
    size = 9 if len(columns) <= 6 else 8
    for j, inches in enumerate(_column_widths(block, frame, columns, width / Inches(1), size)):
        table.columns[j].width = int(Inches(inches))
    for j, name in enumerate(columns):
        _cell(
            table.cell(0, j),
            name,
            size,
            brand.WHITE,
            brand.BLUE,
            bold=True,
            right=_numeric(block, name),
        )
    for i, record in enumerate(frame.itertuples(index=False), start=1):
        fill = brand.PANEL if i % 2 == 0 else brand.WHITE
        for j, value in enumerate(record):
            column = columns[j]
            _cell(
                table.cell(i, j),
                format_value(value, block.formats.get(column, "text")),
                size,
                brand.INK,
                fill,
                right=_numeric(block, column),
            )
    for row in table.rows:
        row.height = row_h
    if len(block.frame) > limit:
        _text(
            slide,
            left,
            top + row_h * rows + Inches(0.02),
            width,
            Inches(0.25),
            f"Top {limit} of {len(block.frame):,} lines; the full list is in the Excel workbook.",
            8,
            brand.MUTED,
            italic=True,
        )


def _column_widths(
    block: TableBlock, frame: pd.DataFrame, columns: list[str], width: float, size: float
) -> list[float]:
    """Column widths in inches: every column gets the room its longest value needs
    (codes and numbers never wrap); when that is more than the table width, only
    the long text columns give up room."""
    per_char = 0.0085 * size  # inches per character, roughly, at this font size
    needs = []
    for column in columns:
        texts = [format_value(v, block.formats.get(column, "text")) for v in frame[column].tolist()]
        words = [len(w) for w in column.split()] + [len(t) for t in texts if " " not in t]
        longest_word = max([*words, 3])
        longest = max([len(t) for t in texts] + [len(column) * 0.6, 3])
        needs.append((min(longest, 48) * per_char + 0.12, longest_word * per_char + 0.12))
    total = sum(n for n, _ in needs)
    if total <= width:
        return [n * width / total for n, _ in needs]
    minimum = sum(m for _, m in needs)
    if minimum >= width:
        return [m * width / minimum for _, m in needs]
    spare = width - minimum
    extra = [n - m for n, m in needs]
    return [m + e * spare / sum(extra) for (_, m), e in zip(needs, extra, strict=True)]


def _cell(
    cell: Any,
    text: str,
    size: float,
    colour: str,
    fill: str,
    bold: bool = False,
    right: bool = False,
) -> None:
    cell.fill.solid()
    cell.fill.fore_color.rgb = _rgb(fill)
    cell.margin_left = cell.margin_right = Inches(0.05)
    cell.margin_top = cell.margin_bottom = Inches(0.02)
    cell.vertical_anchor = MSO_ANCHOR.MIDDLE
    frame = cell.text_frame
    frame.word_wrap = True
    paragraph = frame.paragraphs[0]
    paragraph.alignment = PP_ALIGN.RIGHT if right else PP_ALIGN.LEFT
    run = paragraph.add_run()
    run.text = text
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.name = brand.BODY_FONT
    run.font.color.rgb = _rgb(colour)


def _finite(value: float | None) -> float:
    """Chart values must be numbers: a missing or infinite value is drawn as 0."""
    number = float(value) if value is not None else 0.0
    return number if math.isfinite(number) else 0.0


def number_format(values: Sequence[float], value_format: ValueFormat) -> str:
    """An Excel number format that keeps axis and labels short (1.2M, 3.4B)."""
    if value_format == "percent":
        return "0.0%"
    largest = max((abs(v) for v in values), default=0.0)
    if largest >= 1e9:
        return '#,##0.00,,,"B"'
    if largest >= 1e6:
        return '#,##0.0,,"M"'
    if largest >= 1e5:
        return '#,##0,"K"'
    if value_format in ("money", "integer") or largest >= 100:
        return "#,##0"
    return "#,##0.0"


def _add_chart(slide: Any, chart: ChartBlock, left: int, top: int, width: int, height: int) -> None:
    data = CategoryChartData()  # type: ignore[no-untyped-call]
    data.categories = [str(c) for c in chart.categories]
    values: list[float] = []
    for name, series in chart.series.items():
        floats = [_finite(v) for v in series]
        values += floats
        data.add_series(name, floats)  # type: ignore[no-untyped-call]
    graphic = slide.shapes.add_chart(_CHART_TYPES[chart.kind], left, top, width, height, data)
    plotted = graphic.chart
    plotted.font.size = Pt(9)
    plotted.font.name = brand.BODY_FONT
    plotted.has_title = True
    plotted.chart_title.text_frame.text = chart.title
    title_run = plotted.chart_title.text_frame.paragraphs[0].runs[0]
    title_run.font.size, title_run.font.bold = Pt(12), True
    title_run.font.color.rgb = _rgb(brand.BLUE)
    fmt = number_format(values, chart.value_format)
    plot = plotted.plots[0]
    if chart.kind == "pie":
        for k, point in enumerate(plot.series[0].points):
            point.format.fill.solid()
            point.format.fill.fore_color.rgb = _rgb(brand.SERIES[k % len(brand.SERIES)])
    else:
        for k, series in enumerate(plot.series):
            colour = _rgb(brand.SERIES[k % len(brand.SERIES)])
            if chart.kind == "line":
                series.format.line.color.rgb = colour
                series.format.line.width = Pt(2.25)
                series.smooth = False
            else:
                series.format.fill.solid()
                series.format.fill.fore_color.rgb = colour
        if chart.kind in ("bar", "column"):
            plot.gap_width = 60
        value_axis = plotted.value_axis
        value_axis.tick_labels.number_format = fmt
        value_axis.tick_labels.number_format_is_linked = False
        value_axis.has_major_gridlines = True
        value_axis.major_gridlines.format.line.color.rgb = _rgb(brand.LINE)
        value_axis.format.line.fill.background()
        if chart.kind == "bar":
            plotted.category_axis.reverse_order = True  # first category on top
        if chart.labels and chart.kind != "stacked":
            plot.has_data_labels = True
            labels = plot.data_labels
            labels.number_format = fmt
            labels.number_format_is_linked = False
            labels.font.size = Pt(8)
            labels.font.color.rgb = _rgb(brand.INK)
            if chart.kind in ("bar", "column"):
                labels.position = XL_LABEL_POSITION.OUTSIDE_END
            elif chart.kind == "line":
                labels.position = XL_LABEL_POSITION.ABOVE
    plotted.has_legend = len(chart.series) > 1 or chart.kind == "pie"
    if plotted.has_legend:
        plotted.legend.position = XL_LEGEND_POSITION.BOTTOM
        plotted.legend.include_in_layout = False


def build_pptx(document: ReportDocument, template: Path | None = None) -> bytes:
    if template is None:
        return Deck(document).build()
    prs = Presentation(str(template))
    _remove_existing_slides(prs)
    _template_deck(prs, document)
    buffer = io.BytesIO()
    prs.save(buffer)
    return buffer.getvalue()
