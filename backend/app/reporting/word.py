"""Word reports, on an uploaded organisation template or the built-in RMS layout.

Built-in layout (RMS Brand Guide): a blue cover band with the white logo, the
title in Effra and the tagline; the colour logo in the page header and the
report, period and page number in the footer; Effra headings in corporate blue,
Open Sans body text; KPI cards, highlight and "So what" boxes in the brand
tones; tables with a blue header row; charts drawn in the brand colours.

Template convention (documented in docs/operations.md): the .docx may contain
`{{ title }}`, `{{ subtitle }}` and `{{ report_date }}` anywhere (cover,
header, footer), and must contain a paragraph with only `{{p body }}` where
the report content goes. Everything else in the template (logo, styles,
cover, header, footer) is kept as designed.
"""

import io
from collections.abc import Sequence
from pathlib import Path
from typing import Any

import pandas as pd
from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor
from docxtpl import DocxTemplate

from app.reporting import brand
from app.reporting.charts import chart_png
from app.reporting.documents import (
    Highlight,
    KpiCard,
    ReportDocument,
    TableBlock,
    format_value,
)


class TemplateError(Exception):
    pass


def _rgb(colour: str) -> RGBColor:
    return RGBColor(*brand.rgb(colour))


def _shade(cell: Any, colour: str) -> None:
    properties = cell._tc.get_or_add_tcPr()
    shading = OxmlElement("w:shd")
    shading.set(qn("w:val"), "clear")
    shading.set(qn("w:color"), "auto")
    shading.set(qn("w:fill"), colour)
    properties.append(shading)


def _borders(table: Any, colour: str = brand.LINE, size: int = 4, inside: bool = True) -> None:
    properties = table._tbl.tblPr
    borders = OxmlElement("w:tblBorders")
    edges = ("top", "left", "bottom", "right", *(("insideH", "insideV") if inside else ()))
    for edge in edges:
        element = OxmlElement(f"w:{edge}")
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), str(size))
        element.set(qn("w:color"), colour)
        borders.append(element)
    properties.append(borders)


def _left_bar(cell: Any, colour: str) -> None:
    properties = cell._tc.get_or_add_tcPr()
    borders = OxmlElement("w:tcBorders")
    left = OxmlElement("w:left")
    left.set(qn("w:val"), "single")
    left.set(qn("w:sz"), "24")
    left.set(qn("w:color"), colour)
    borders.append(left)
    properties.append(borders)


def _run(
    paragraph: Any,
    text: str,
    size: float = 10,
    colour: str = brand.INK,
    bold: bool = False,
    italic: bool = False,
    font: str = brand.BODY_FONT,
) -> Any:
    run = paragraph.add_run(text)
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic
    run.font.name = font
    run._element.rPr.rFonts.set(qn("w:eastAsia"), font)
    run.font.color.rgb = _rgb(colour)
    return run


def _page_field(paragraph: Any) -> None:
    run = paragraph.add_run()
    for kind, text in (("begin", None), (None, "PAGE"), ("end", None)):
        if kind:
            element = OxmlElement("w:fldChar")
            element.set(qn("w:fldCharType"), kind)
        else:
            element = OxmlElement("w:instrText")
            element.set(qn("xml:space"), "preserve")
            element.text = text
        run._r.append(element)
    run.font.size = Pt(8)


def _heading(container: Any, text: str, level: int = 1) -> None:
    paragraph = container.add_paragraph()
    paragraph.paragraph_format.space_before = Pt(14 if level == 1 else 8)
    paragraph.paragraph_format.space_after = Pt(4)
    paragraph.paragraph_format.keep_with_next = True
    _run(paragraph, text, 16 if level == 1 else 12, brand.BLUE, bold=True, font=brand.HEADING_FONT)
    if level == 1:
        bar = container.add_paragraph()
        bar.paragraph_format.space_after = Pt(6)
        _run(bar, "▬▬", 8, brand.YELLOW)


def _add_table(container: Any, block: TableBlock) -> None:
    frame = block.frame if block.max_rows is None else block.frame.head(block.max_rows)
    columns = [str(c) for c in frame.columns]
    if not columns:
        return
    numeric = {
        c for c in columns if block.formats.get(c) in ("money", "number", "integer", "percent")
    }
    table = container.add_table(rows=1, cols=len(columns))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    _borders(table)
    size = 8.5 if len(columns) > 6 else 9
    for j, name in enumerate(columns):
        cell = table.rows[0].cells[j]
        _shade(cell, brand.BLUE)
        paragraph = cell.paragraphs[0]
        if name in numeric:
            paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        _run(paragraph, name, size, brand.WHITE, bold=True)
    for i, record in enumerate(frame.itertuples(index=False)):
        cells = table.add_row().cells
        for j, value in enumerate(record):
            if i % 2:
                _shade(cells[j], brand.PANEL)
            paragraph = cells[j].paragraphs[0]
            if columns[j] in numeric:
                paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
            _run(paragraph, format_value(value, block.formats.get(columns[j], "text")), size)
    if block.total_row and len(frame):
        cells = table.add_row().cells
        for j, name in enumerate(columns):
            _shade(cells[j], brand.tint(brand.BLUE, 0.15))
            paragraph = cells[j].paragraphs[0]
            text = ""
            if j == 0:
                text = "Total"
            elif block.formats.get(name) in ("money", "number", "integer"):
                text = format_value(block.frame[name].sum(), block.formats[name])
                paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
            _run(paragraph, text, size, bold=True)
    if block.max_rows is not None and len(block.frame) > block.max_rows:
        note = container.add_paragraph()
        _run(
            note,
            f"Showing the first {block.max_rows} of {len(block.frame)} lines; "
            "the full list is in the Excel workbook.",
            8,
            brand.MUTED,
            italic=True,
        )


def _caption(container: Any, text: str) -> None:
    paragraph = container.add_paragraph()
    paragraph.paragraph_format.space_before = Pt(8)
    paragraph.paragraph_format.keep_with_next = True
    _run(paragraph, text, 10.5, brand.BLUE, bold=True, font=brand.HEADING_FONT)


def _cards(container: Any, cards: Sequence[KpiCard]) -> None:
    cards = list(cards)[:6]
    if not cards:
        return
    per_row = 3 if len(cards) > 4 else len(cards)
    rows = (len(cards) + per_row - 1) // per_row
    table = container.add_table(rows=rows, cols=per_row)
    _borders(table, brand.WHITE, 12)
    for k, card in enumerate(cards):
        cell = table.rows[k // per_row].cells[k % per_row]
        accent, fill = brand.TONES[card.tone]
        _shade(cell, fill)
        value = cell.paragraphs[0]
        value.alignment = WD_ALIGN_PARAGRAPH.CENTER
        _run(value, card.value, 18, accent, bold=True, font=brand.HEADING_FONT)
        label = cell.add_paragraph()
        label.alignment = WD_ALIGN_PARAGRAPH.CENTER
        _run(label, card.label, 8.5, brand.INK, bold=True)
        if card.note:
            note = cell.add_paragraph()
            note.alignment = WD_ALIGN_PARAGRAPH.CENTER
            _run(note, card.note, 7.5, brand.MUTED)
    container.add_paragraph()


def _boxes(container: Any, items: Sequence[Highlight], per_row: int = 1) -> None:
    items = list(items)
    if not items:
        return
    rows = (len(items) + per_row - 1) // per_row
    table = container.add_table(rows=rows, cols=per_row)
    _borders(table, brand.WHITE, 12)
    for k, item in enumerate(items):
        cell = table.rows[k // per_row].cells[k % per_row]
        accent, fill = brand.TONES[item.tone]
        _shade(cell, fill)
        _left_bar(cell, accent)
        _run(cell.paragraphs[0], item.title, 10, accent, bold=True, font=brand.HEADING_FONT)
        for line in item.text.split("\n"):
            _run(cell.add_paragraph(), line, 9)
    container.add_paragraph()


def _chart(container: Any, chart: Any) -> None:
    paragraph = container.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.add_run().add_picture(io.BytesIO(chart_png(chart)), width=Inches(6.3))


def _fill(container: Any, document: ReportDocument) -> None:
    if document.summary_message or document.kpis or document.highlights:
        _heading(container, "Executive summary")
        if document.summary_message:
            _run(container.add_paragraph(), document.summary_message, 11, brand.INK, bold=True)
        _cards(container, document.kpis)
        _boxes(container, document.highlights, per_row=1)
    for section in document.sections:
        _heading(container, section.heading)
        if section.message:
            message = container.add_paragraph()
            message.paragraph_format.keep_with_next = True
            _run(message, section.message, 10, brand.MUTED, italic=True)
        _cards(container, section.kpis)
        for paragraph in section.paragraphs:
            _run(container.add_paragraph(), paragraph)
        for bullet in section.bullets:
            _run(container.add_paragraph(style="List Bullet"), bullet)
        for chart in section.charts:
            _chart(container, chart)
        for block in section.tables:
            _caption(container, block.title)
            _add_table(container, block)
        _boxes(container, section.highlights, per_row=2 if len(section.highlights) > 1 else 1)
        if section.takeaway:
            _boxes(container, [Highlight("So what", section.takeaway, "blue")])
    if document.recommendations:
        _heading(container, "Recommendations and next steps")
        frame = pd.DataFrame(
            {
                "No.": [str(k) for k in range(1, len(document.recommendations) + 1)],
                "Action": [r.action for r in document.recommendations],
                "Proposed owner": [r.owner for r in document.recommendations],
                "Timing": [r.timing for r in document.recommendations],
            }
        )
        _add_table(container, TableBlock("Recommendations", frame))
        widths = (Cm(1.0), Cm(9.5), Cm(4.0), Cm(2.5))
        for row in container.tables[-1].rows:
            for cell, width in zip(row.cells, widths, strict=True):
                cell.width = width
    if document.notes:
        _heading(container, "Data notes")
        _boxes(container, document.notes, per_row=2 if len(document.notes) > 1 else 1)
    if document.checks:
        _heading(container, "Report checks")
        checks = TableBlock(
            "Checks",
            frame=pd.DataFrame(
                [
                    {
                        "Check": c.label,
                        "Expected": c.expected,
                        "Actual": c.actual,
                        "Result": "Passed" if c.passed else "Failed",
                    }
                    for c in document.checks
                ]
            ),
            formats={"Expected": "number", "Actual": "number"},
        )
        _add_table(container, checks)


def _cover(doc: Any, document: ReportDocument) -> None:
    band = doc.add_table(rows=1, cols=1)
    cell = band.rows[0].cells[0]
    _shade(cell, brand.BLUE)
    first = cell.paragraphs[0]
    first.paragraph_format.space_before = Pt(10)
    content, size = brand.logo("white"), brand.logo_size("white")
    if content and size:
        first.add_run().add_picture(io.BytesIO(content), height=Cm(1.6))
    title = cell.add_paragraph()
    title.paragraph_format.space_before = Pt(18)
    _run(title, document.title.upper(), 24, brand.WHITE, bold=True, font=brand.HEADING_FONT)
    _run(cell.add_paragraph(), document.subtitle, 12, brand.WHITE, italic=True)
    _run(cell.add_paragraph(), document.date_text, 11, brand.WHITE)
    if document.audience or document.draft:
        line = " · ".join(p for p in (document.audience, "DRAFT" if document.draft else "") if p)
        _run(cell.add_paragraph(), line, 9, brand.WHITE)
    tagline = cell.add_paragraph()
    tagline.paragraph_format.space_before = Pt(14)
    tagline.paragraph_format.space_after = Pt(10)
    _run(tagline, brand.TAGLINE, 11, brand.YELLOW, bold=True, font=brand.HEADING_FONT)


def _header_footer(doc: Any, document: ReportDocument) -> None:
    section = doc.sections[0]
    header = section.header.paragraphs[0]
    header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    content = brand.logo("colour")
    if content:
        header.add_run().add_picture(io.BytesIO(content), height=Cm(0.9))
    footer = section.footer.paragraphs[0]
    _run(
        footer,
        brand.footer(document.title, document.date_text, document.draft) + "   ·   Page ",
        8,
        brand.MUTED,
    )
    _page_field(footer)


def build_docx(document: ReportDocument, template: Path | None = None) -> bytes:
    buffer = io.BytesIO()
    if template is None:
        doc = Document()
        normal = doc.styles["Normal"]
        normal.font.name = brand.BODY_FONT
        normal.element.rPr.rFonts.set(qn("w:eastAsia"), brand.BODY_FONT)
        normal.font.size = Pt(10)
        normal.font.color.rgb = _rgb(brand.INK)
        section = doc.sections[0]
        section.orientation = WD_ORIENT.PORTRAIT
        section.page_width, section.page_height = Cm(21.0), Cm(29.7)
        for side in ("left_margin", "right_margin"):
            setattr(section, side, Cm(2.0))
        section.top_margin, section.bottom_margin = Cm(2.2), Cm(2.0)
        _header_footer(doc, document)
        _cover(doc, document)
        doc.add_paragraph()
        _fill(doc, document)
        doc.save(buffer)
        return buffer.getvalue()

    tpl = DocxTemplate(str(template))
    variables = tpl.get_undeclared_template_variables()
    if "body" not in variables:
        raise TemplateError(
            "The Word template has no {{p body }} placeholder for the report content."
        )
    body = tpl.new_subdoc()
    _fill(body, document)
    tpl.render(
        {
            "title": document.title,
            "subtitle": document.subtitle,
            "report_date": document.date_text,
            "body": body,
        }
    )
    tpl.save(buffer)
    return buffer.getvalue()
